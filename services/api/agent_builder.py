"""Agent Builder API: brief → agent → edit → test on Retell → publish.

    GET  /agent-builder/options                     voices, choices, what is configured
    POST /agent-builder/drafts                      a brief → a new agent (server-sent events)
    GET  /agent-builder/agents                      this organization's agents
    GET  /agent-builder/agents/{agent_id}           one agent, with its readiness and call config
    PUT  /agent-builder/agents/{agent_id}           save the review screen
    POST /agent-builder/agents/{agent_id}/revise    "Ask Tara to change something"
    POST /agent-builder/agents/{agent_id}/questions/generate   three more questions
    POST /agent-builder/agents/{agent_id}/test-chat            a typed rehearsal, on the real prompt
    POST /agent-builder/agents/{agent_id}/test-call            a Retell web call → access token
    POST /agent-builder/agents/{agent_id}/score                score a test against the rubric
    POST /agent-builder/agents/{agent_id}/publish              freeze a version
    POST /agent-builder/agents/{agent_id}/invites              an access code + link for one participant
    GET  /agent-builder/agents/{agent_id}/sessions             who took it, and how they scored

    POST /api/agent-builder/retell-webhook          Retell's call events (signed; public)

Mounted under /api/recruiter behind `builder_scope`, like the Scenario Builder:
the recruiter guard, plus the local-prototype exception that is off by default.
Responses carry the rubric, which is the point of an authoring surface. The
Retell payload never does, and `leaked_cues` refuses any payload that would.

There is ONE Retell agent (`RETELL_AGENT_BUILDER_AGENT_ID`). Nothing here creates
another. A built agent reaches it as per-call variables and `agent_override`.
"""
from __future__ import annotations

import copy
import hashlib
import hmac
import json
import re
import time
from typing import Any

from fastapi import APIRouter, BackgroundTasks, Depends, HTTPException, Request, Response
from fastapi.responses import StreamingResponse
from pydantic import BaseModel, Field

from services import config
from services.ai.brain import LLMError
from services.ai.workloads import agent_builder as ab
from services.assessment import agent_builder_retell as rx
from services.evaluation import simulation as sim
from services.notify import email
from services.notify import template as tmpl
from services.data import built_agents as store
from services.security import ratelimit

router = APIRouter(prefix="/agent-builder", tags=["agent-builder"])
#: Retell calls this one. It has no principal: its proof is the signature.
webhook_router = APIRouter(tags=["agent-builder"])

RETELL = "https://api.retellai.com"


# --------------------------------------------------------------------------- #
#  Ownership
# --------------------------------------------------------------------------- #
def _org(request: Request) -> str:
    """The caller's organization. "local" only under the local-prototype
    exception, where there is no principal at all."""
    from services.security import principal as security

    p = security.optional_principal(request)
    return p.organization_id if p else "local"


LOCKED = ("This role-play is published, so it's locked: you can test it, invite people and read its reports, "
          "but not change it. Duplicate it to make an editable copy.")


def _editable(agent_id: str, request: Request) -> dict[str, Any]:
    """The agent, if it may still be changed. A published role-play is locked:
    every result already collected must stay tied to exactly what people took."""
    row = _row(agent_id, request)
    if row.get("published"):
        raise HTTPException(409, {"error": "locked", "message": LOCKED})
    return row


def _row(agent_id: str, request: Request) -> dict[str, Any]:
    row = store.load(agent_id)
    # Another organization's agent is indistinguishable from one that does not
    # exist, so an id cannot be probed for.
    if row is None or row.get("org_id") != _org(request):
        raise HTTPException(404, "No such agent.")
    return row


# --------------------------------------------------------------------------- #
#  Shape
# --------------------------------------------------------------------------- #
def _s(v: Any) -> str:
    return "" if v is None else str(v).strip()


def _one(v: Any, options: tuple[str, ...], default: str) -> str:
    return v if v in options else default


def clean_cfg(cfg: dict[str, Any], base: dict[str, Any] | None = None) -> dict[str, Any]:
    base = base or {}
    cfg = {**base, **(cfg or {})}
    v = cfg.get("voice") if cfg.get("voice") in rx.VOICE_KEYS else rx.DEFAULT_VOICE
    return {
        "tone": _one(cfg.get("tone"), rx.DIFFICULTIES, "Realistic"),
        "depth": _one(cfg.get("depth"), rx.DEPTHS, "Probing"),
        "followups": bool(cfg.get("followups", True)),
        # Every agent is a voice conversation: no video, no typed chat.
        "format": rx.FORMATS[0],
        # The persona always opens, with no introduction sound: both settings
        # were removed from the builder as choices nobody needed to make.
        "speaker": rx.SPEAKERS[0],
        "sound": rx.SOUNDS[0],
        "ending": _one(cfg.get("ending"), rx.ENDINGS, rx.ENDINGS[0]),
        "voice": v,
        # The language follows the voice. Voice is chosen in Persona, and a
        # Hindi voice told to speak French would be a configuration nobody meant.
        "language": rx.voice(v).language,
        "consent": bool(cfg.get("consent", True)),
        # Defaults for the proctoring suite, which applies them; the builder only stores them.
        "image_proctoring": bool(cfg.get("image_proctoring", False)),
        "safe_browser": bool(cfg.get("safe_browser", False)),
        **_purpose_cfg(cfg),
    }


def _purpose_cfg(cfg: dict[str, Any]) -> dict[str, Any]:
    """The attempt policy. (Purpose was removed from the builder; every role-play is General.)"""
    return {"attempts": cfg.get("attempts") if cfg.get("attempts") in sim.ATTEMPTS else "1"}


def persona_only(agent: dict[str, Any]) -> dict[str, Any]:
    """Participants only meet the persona, so no line says "Tara".

    Whatever wrote the text (a draft, Ask Tara, Generate more, the Skill
    Master, or a person typing), every "Tara" in the title, description,
    instructions, opening, closing and questions becomes the persona's
    first name.
    """
    first = rx.persona_first((agent.get("persona") or {}).get("name"))
    fix = lambda t: no_interview(rx._TARA.sub(first, t or ""))  # noqa: E731
    for key in ("title", "type_label", "description", "instructions", "opening_line", "closing_line",
                "context", "additional_context"):
        agent[key] = fix(agent.get(key))
    for q in agent.get("questions") or []:
        q["text"] = fix(q.get("text"))
    return agent


no_interview = rx.no_interview


CONTEXT_MAX = ab.BRIEF_MAX
ADDITIONAL_MAX = 3000


def _text(v: Any) -> str:
    """Long text kept as written: line breaks matter in a brief, only the ends are trimmed."""
    return str(v or "").replace("\r\n", "\n").strip() if isinstance(v, (str, int, float)) else ""


def clean_agent(agent: dict[str, Any]) -> dict[str, Any]:
    return persona_only(_clean_agent(agent))


def _clean_agent(agent: dict[str, Any]) -> dict[str, Any]:
    persona = agent.get("persona") or {}
    rubric = []
    for r in agent.get("rubric") or []:
        try:
            w = max(0, min(100, int(round(float(r.get("weight") or 0)))))
        except (TypeError, ValueError):
            w = 0
        rubric.append({"name": _s(r.get("name"))[:120], "anchor": _s(r.get("anchor"))[:600], "weight": w})
    return {
        "title": _s(agent.get("title"))[:120] or "New agent",
        "type_label": _s(agent.get("type_label"))[:80],
        "persona": {k: _s(persona.get(k))[:160] for k in ("name", "role", "style")},
        "description": _s(agent.get("description"))[:2000],
        "instructions": _s(agent.get("instructions"))[:6000],
        # The one-pager's limits: Scenario Context 10,000, Additional Context 3,000.
        # Evaluation guidance is the scorer's alone and never reaches the voice agent.
        "context": rx.without_lines(_text(agent.get("context"))[:CONTEXT_MAX], _text(agent.get("evaluation_context"))),
        "additional_context": _text(agent.get("additional_context"))[:ADDITIONAL_MAX],
        "evaluation_context": _text(agent.get("evaluation_context"))[:CONTEXT_MAX],
        "opening_line": _s(agent.get("opening_line"))[:600],
        "closing_line": rx.closing_without_question(_s(agent.get("closing_line"))[:600]) if _s(agent.get("closing_line")) else "",
        # A question may test several skills (`tags`); `tag` stays as the first for older readers.
        "questions": [ab.with_tags(_s(q.get("text"))[:600], ab.question_tags(q, {r["name"].lower(): r["name"] for r in rubric}))
                      for q in (agent.get("questions") or [])[:30] if isinstance(q, dict) and _s(q.get("text"))],
        "rubric": rubric[:10],
        "exhibits": clean_exhibits(agent.get("exhibits")),
        "depth": _one(agent.get("depth"), rx.DEPTHS, "Probing"),
        "voice": agent.get("voice") if agent.get("voice") in rx.VOICE_KEYS else rx.DEFAULT_VOICE,
    }


CHART_TYPES = ("bar", "line", "pie")
_FILE_ID = re.compile(r"^[a-f0-9]{24}\.(png|jpg|webp)$")


def clean_exhibits(raw: Any) -> list[dict[str, Any]]:
    """Up to six exhibits the participant sees during the conversation: a chart
    from a small table, or an uploaded image, each with what the persona knows about it."""
    out = []
    for e in (raw if isinstance(raw, list) else [])[:6]:
        if not isinstance(e, dict):
            continue
        kind = "image" if e.get("kind") == "image" else "chart"
        item = {"kind": kind, "title": _s(e.get("title"))[:80] or f"Exhibit {len(out) + 1}",
                "description": _s(e.get("description"))[:600]}
        if kind == "chart":
            c = e.get("chart") if isinstance(e.get("chart"), dict) else {}
            labels = [_s(x)[:30] for x in (c.get("labels") or [])][:12]
            vals = []
            for v in (c.get("values") or [])[:len(labels)]:
                try:
                    vals.append(round(float(v), 4))
                except (TypeError, ValueError):
                    vals.append(0.0)
            vals += [0.0] * (len(labels) - len(vals))
            item["chart"] = {"type": c.get("type") if c.get("type") in CHART_TYPES else "bar",
                             "labels": labels, "values": vals, "unit": _s(c.get("unit"))[:12]}
        else:
            f = _s(e.get("file"))
            if not _FILE_ID.match(f):
                continue
            item["file"] = f
        out.append(item)
    return out


def clean_fields(fields: dict[str, Any], cfg: dict[str, Any]) -> dict[str, Any]:
    out = {k: _s((fields or {}).get(k))[:400] for k in ("type", "role", "persona", "skills")}
    # Assessment is on hold: every agent is a Role-play, whatever an older
    # draft, Ask Tara or a request says.
    out["type"] = "Role-play"
    out["difficulty"] = cfg["tone"]
    out["length"] = f"{rx.lengths(cfg)[0]} min"
    return out


def readiness(row: dict[str, Any]) -> list[dict[str, Any]]:
    a, f = row["agent"], row["fields"]
    total = sum(r["weight"] for r in a["rubric"])
    return [
        {"id": "scenario", "label": "Scenario details", "done": all(f.get(k) for k in ("type", "role", "skills"))},
        {"id": "desc", "label": "Description", "done": bool(a["description"])},
        {"id": "context", "label": "Instructions", "done": bool(a["instructions"])},
        {"id": "persona", "label": "Persona", "done": bool(a["persona"]["name"])},
        {"id": "rubric", "label": "Skills", "done": bool(row.get("reviewed")) and total == 100 and len(a["rubric"]) > 0},
    ]


def call_body(row: dict[str, Any]) -> dict[str, Any]:
    webhook = f"{config.PUBLIC_URL}/api/agent-builder/retell-webhook" if config.PUBLIC_URL else ""
    return rx.web_call_body(row, config.RETELL_AGENT_BUILDER_AGENT_ID or "RETELL_AGENT_BUILDER_AGENT_ID", webhook)


def public(row: dict[str, Any]) -> dict[str, Any]:
    body = call_body(row)
    target, cap = rx.lengths(row["cfg"])
    return {
        "agent_id": row["agent_id"],
        "brief": row.get("brief", ""),
        "mode": "roleplay",   # Assessment is on hold
        "fields": row["fields"],
        # Shown with today's wording rules ("conversation", the persona's name), even
        # for an agent saved or published before them; edits save it this way too.
        "agent": persona_only(copy.deepcopy(row["agent"])),
        "cfg": row["cfg"],
        "reviewed": bool(row.get("reviewed")),
        "version": int(row.get("version") or 0),
        "published_version": int((row.get("published") or {}).get("version") or 0),
        "locked": bool(row.get("published")),
        "readiness": readiness(row),
        "length": {"target_minutes": target, "cap_minutes": cap},
        "call_config": body,
        "leaked_cues": rx.leaked_cues(body, row),
        "tests": (row.get("tests") or [])[-10:],
        "candidate_link": next(({"code": i["code"], "path": f"/participant?code={i['code']}"}
                                for i in row.get("invites") or [] if i.get("shared") and not i.get("revoked")), None),
        "open_link": _open_link(row),
        "updated_at": row.get("updated_at", 0),
    }


def _open_link(row: dict[str, Any]) -> dict[str, Any]:
    inv = next((i for i in row.get("invites") or [] if i.get("shared")), None)
    on = bool(inv and not inv.get("revoked"))
    return {"enabled": on, "code": inv["code"] if on else "",
            "path": f"/participant?code={inv['code']}" if on else "",
            "proctoring": (inv or {}).get("proctoring") or _proctoring(row)}


def _proctoring(row: dict[str, Any], image: bool | None = None, safe: bool | None = None,
                current: dict[str, Any] | None = None) -> dict[str, bool]:
    """Proctoring for one invitation or the open link: what was asked, else what it had, else the role-play default."""
    cfg = (row.get("published") or {}).get("cfg") or row["cfg"]
    cur = current or {}
    return {"image_proctoring": bool(image if image is not None else cur.get("image_proctoring", cfg.get("image_proctoring", False))),
            "safe_browser": bool(safe if safe is not None else cur.get("safe_browser", cfg.get("safe_browser", False)))}


def _new_row(p: dict[str, Any], c: dict[str, Any], brief: str, mode: str, org: str,
             contexts: dict[str, str] | None = None) -> dict[str, Any]:
    ctx = contexts or {}
    agent = clean_agent({**p["agent"], **c, "context": ctx.get("persona_context", ""),
                         "evaluation_context": ctx.get("evaluation_context", "")})
    cfg = clean_cfg({"tone": p["fields"]["difficulty"], "depth": agent["depth"], "voice": agent["voice"]})
    return {
        "agent_id": store.new_id(agent["title"]),
        "org_id": org,
        "brief": brief,
        "mode": mode,
        "fields": clean_fields(p["fields"], cfg),
        "agent": agent,
        "cfg": cfg,
        "reviewed": False,
        "version": 0,
        "tests": [],
    }


# --------------------------------------------------------------------------- #
#  Options
# --------------------------------------------------------------------------- #
@router.get("/options")
def options() -> dict[str, Any]:
    from services.ai.gateway import get_gateway

    return {
        "voices": [{"key": v.key, "label": v.label, "language": v.language, "locale": v.locale}
                   for v in rx.VOICES],
        "difficulties": rx.DIFFICULTIES, "depths": rx.DEPTHS, "depth_minutes": rx.DEPTH_MINUTES,
        "endings": rx.ENDINGS, "speakers": rx.SPEAKERS, "formats": rx.FORMATS, "sounds": rx.SOUNDS,
        "languages": sorted({v.language for v in rx.VOICES}),
        "test_call_minutes": rx.TEST_CALL_MINUTES,
        "email_configured": email.configured(), "email_from": email.sender(),
        "support_url": config.SUPPORT_URL,
        "model_configured": get_gateway().live,
        "voice_configured": bool(config.RETELL_API_KEY and config.RETELL_AGENT_BUILDER_AGENT_ID),
    }


# --------------------------------------------------------------------------- #
#  Draft
# --------------------------------------------------------------------------- #
class DraftBody(BaseModel):
    brief: str = Field(min_length=1, max_length=ab.BRIEF_MAX)
    mode: str = "roleplay"


#: Briefs longer than this carry real detail (situations, expected answers) and are split.
SPLIT_FROM = 1200


def _sse(event: str, data: Any) -> str:
    return f"event: {event}\ndata: {json.dumps(data)}\n\n"


@router.post("/drafts")
def draft(
    body: DraftBody,
    request: Request,
    _: None = Depends(ratelimit.limiter("generation")),
) -> StreamingResponse:
    """Build an agent from a brief, streamed in two real stages.

    `plan` arrives once the scenario, persona and lines exist; `content` once
    the questions and rubric do; `done` carries the saved agent. Screen 2 shows
    each as it lands.
    """
    brief = body.brief.strip()
    if not brief:
        raise HTTPException(422, "Describe the agent you want to build.")
    mode = "roleplay"   # Assessment is on hold
    org = _org(request)

    def events():
        try:
            # A short brief is all scenario; a long one is split so the persona
            # never sees the answer key (expected answers, accuracy rules).
            contexts, by0 = (ab.split_context(brief) if len(brief) > SPLIT_FROM
                             else ({"persona_context": "", "evaluation_context": ""}, "none"))
            yield _sse("context", {"persona_chars": len(contexts["persona_context"]),
                                   "evaluation_chars": len(contexts["evaluation_context"]), "drafted_by": by0})
            p, by1 = ab.plan(brief, mode)
            yield _sse("plan", {**p, "drafted_by": by1})
            c, by2 = ab.content(p, contexts)
            yield _sse("content", {**c, "drafted_by": by2})
            row = store.save(_new_row(p, c, brief, mode, org, contexts))
            yield _sse("done", public(row))
        except Exception as exc:  # noqa: BLE001 — the stream must end with a reason
            yield _sse("error", {"message": "Tara couldn't finish the draft. Try again.",
                                 "detail": type(exc).__name__})

    return StreamingResponse(events(), media_type="text/event-stream",
                             headers={"Cache-Control": "no-store", "X-Accel-Buffering": "no"})


# --------------------------------------------------------------------------- #
#  Read and save
# --------------------------------------------------------------------------- #
@router.get("/agents")
def list_agents(request: Request) -> dict[str, Any]:
    org = _org(request)
    return {"agents": [
        {"agent_id": r["agent_id"], "title": no_interview(r["agent"]["title"]), "type_label": no_interview(r["agent"]["type_label"]),
         "persona": ((r["agent"].get("persona") or {}).get("name") or ""),
         "description": no_interview(r["agent"].get("description") or ""),
         "minutes": rx.DEPTH_MINUTES.get((r.get("cfg") or {}).get("depth") or "", 0),
         "published_version": int((r.get("published") or {}).get("version") or 0),
         "updated_at": r.get("updated_at", 0)}
        for r in sorted(store.list_all(), key=lambda r: r.get("updated_at", 0), reverse=True)
        if r.get("org_id") == org
    ]}


@router.get("/agents/{agent_id}")
def get_agent(agent_id: str, request: Request) -> dict[str, Any]:
    return public(_row(agent_id, request))


class SaveBody(BaseModel):
    fields: dict[str, Any]
    agent: dict[str, Any]
    cfg: dict[str, Any]
    reviewed: bool = False


@router.put("/agents/{agent_id}")
def save_agent(agent_id: str, body: SaveBody, request: Request) -> dict[str, Any]:
    row = _editable(agent_id, request)
    row["agent"] = clean_agent(body.agent)
    row["cfg"] = clean_cfg(body.cfg)
    row["agent"]["depth"], row["agent"]["voice"] = row["cfg"]["depth"], row["cfg"]["voice"]
    row["fields"] = clean_fields(body.fields, row["cfg"])
    row["reviewed"] = body.reviewed
    row["version"] = int(row.get("version") or 0) + 1
    return public(store.save(row))


# --------------------------------------------------------------------------- #
#  Ask Tara, and more questions
# --------------------------------------------------------------------------- #
class ReviseBody(BaseModel):
    instruction: str = Field(min_length=1, max_length=1500)


def _no_model() -> HTTPException:
    return HTTPException(503, "No model is configured, so Tara can't do this. Set OPENROUTER_API_KEY.")


@router.post("/agents/{agent_id}/revise")
def revise(agent_id: str, body: ReviseBody, request: Request,
           _: None = Depends(ratelimit.limiter("generation"))) -> dict[str, Any]:
    row = _editable(agent_id, request)
    try:
        out = ab.revise(row, body.instruction.strip())
    except LLMError as exc:
        raise _no_model() from exc
    # A revision that renames the persona renames it everywhere it is said.
    old_name = (row["agent"].get("persona") or {}).get("name") or ""
    new_name = ((out["agent"].get("persona") or {}).get("name") or "").strip()
    if new_name and new_name != old_name:
        ab.rename_persona(out["agent"], old_name, new_name)
    row["cfg"] = clean_cfg(out["cfg"])
    row["agent"] = clean_agent({**out["agent"], "depth": row["cfg"]["depth"], "voice": row["cfg"]["voice"]})
    row["fields"] = clean_fields(out["fields"], row["cfg"])
    # A changed rubric is an unreviewed rubric.
    row["reviewed"] = False
    row["version"] = int(row.get("version") or 0) + 1
    return {**public(store.save(row)), "summary": out["summary"]}


@router.get("/skill-master")
def skill_master_catalogue() -> dict[str, Any]:
    """The Skill Master picker: domains by category, each with readable skills."""
    from services.data import skill_master

    lib = json.loads((config.ROOT_DIR / "content" / "skill_library.json").read_text())["skills"]
    cat = skill_master.catalogue()
    return {"version": cat["version"], "categories": cat["categories"], "domains": [
        {"id": d["id"], "label": d["label"], "category": d["category"],
         "description": d.get("description", ""), "skills": lib.get(d["id"], [])}
        for d in cat["domains"]
    ]}


class SkillNames(BaseModel):
    names: list[str] = Field(default_factory=list, max_length=12)


@router.post("/agents/{agent_id}/skills/draft")
def draft_skills(agent_id: str, body: SkillNames, request: Request,
                 _: None = Depends(ratelimit.limiter("generation"))) -> dict[str, Any]:
    """Anchors and questions for skills just added from the Skill Master."""
    row = _editable(agent_id, request)
    have = {r["name"].lower(): r for r in row["agent"]["rubric"]}
    names = [n.strip() for n in body.names if n.strip().lower() in have][:12]
    if not names:
        return {**public(row), "added": 0}
    try:
        out = ab.skill_details(row, names)
    except LLMError as exc:
        raise _no_model() from exc
    for name, anchor in out["anchors"].items():
        r = have[name.lower()]
        if not (r.get("anchor") or "").strip():
            r["anchor"] = anchor
    row["agent"]["questions"] += out["questions"]
    persona_only(row["agent"])
    row["reviewed"] = False
    row["version"] = int(row.get("version") or 0) + 1
    return {**public(store.save(row)), "added": len(out["questions"])}


@router.post("/agents/{agent_id}/questions/generate")
def generate_questions(agent_id: str, request: Request,
                       _: None = Depends(ratelimit.limiter("generation"))) -> dict[str, Any]:
    row = _editable(agent_id, request)
    try:
        added = ab.more_questions(row)
    except LLMError as exc:
        raise _no_model() from exc
    row["agent"]["questions"] += added
    persona_only(row["agent"])
    row["version"] = int(row.get("version") or 0) + 1
    return {**public(store.save(row)), "added": len(added)}


# --------------------------------------------------------------------------- #
#  Testing
# --------------------------------------------------------------------------- #
class Turn(BaseModel):
    role: str = "user"          # "agent" | "user"
    text: str = Field(default="", max_length=4000)


class ChatBody(BaseModel):
    messages: list[Turn] = Field(default_factory=list, max_length=80)


@router.post("/agents/{agent_id}/test-chat")
def test_chat(agent_id: str, body: ChatBody, request: Request) -> dict[str, Any]:
    row = _row(agent_id, request)
    try:
        reply = ab.test_reply(row, [m.model_dump() for m in body.messages])
    except LLMError as exc:
        raise _no_model() from exc
    return {"reply": reply}


@router.post("/agents/{agent_id}/test-call")
async def test_call(agent_id: str, request: Request) -> dict[str, Any]:
    """A Retell web call on the one agent, configured for this agent."""
    row = _row(agent_id, request)
    if not (config.RETELL_API_KEY and config.RETELL_AGENT_BUILDER_AGENT_ID):
        raise HTTPException(503, "Voice testing isn't configured. Set RETELL_API_KEY and "
                                 "RETELL_AGENT_BUILDER_AGENT_ID.")
    body = call_body(row)
    leaks = rx.leaked_cues(body, row)
    if leaks:
        # Refuse rather than strip: a payload carrying the rubric means the
        # instructions restate it, and the recruiter should see where.
        raise HTTPException(422, {"message": "The instructions or questions repeat a skill description, "
                                             "so this can't be sent to the voice agent.", "leaked": leaks})
    limit_test(body)
    from services.assessment import slots

    running, limit = slots.retell_concurrency()
    if limit and running >= limit:
        raise HTTPException(503, f"All {limit} voice lines are in use right now. Try again in a few minutes.")
    import httpx

    try:
        async with httpx.AsyncClient(timeout=20) as http:
            r = await http.post(f"{RETELL}/v3/create-web-call", json=body,
                                headers={"Authorization": f"Bearer {config.RETELL_API_KEY}"})
    except Exception as exc:  # noqa: BLE001
        raise HTTPException(502, "Couldn't reach the voice service.") from exc
    if r.status_code >= 400:
        raise HTTPException(502, rx.refusal(r.status_code, r.text)[1])
    out = r.json()
    call_id = str(out.get("call_id") or "")
    row.setdefault("tests", []).append({"call_id": call_id, "kind": "voice", "at": time.time(),
                                        "version": int(row.get("version") or 0)})
    store.save(row)
    return {**{k: out[k] for k in ("call_id", "access_token", "transport", "url", "ice_servers", "expires_at") if k in out}, "call_id": call_id,
            "max_minutes": rx.TEST_CALL_MINUTES}


def limit_test(body: dict[str, Any]) -> dict[str, Any]:
    """A builder test is the real agent, kept short: Retell ends it at the cap,
    and the persona is told the shorter time so it wraps up first instead of
    being cut off mid-sentence."""
    m = rx.TEST_CALL_MINUTES
    body.setdefault("agent_override", {}).setdefault("agent", {})["max_call_duration_ms"] = m * 60_000
    v = body.setdefault("retell_llm_dynamic_variables", {})
    v["target_minutes"], v["max_minutes"] = str(max(1, m - 1)), str(m)
    return body


class ScoreBody(BaseModel):
    transcript: list[Turn] = Field(default_factory=list, max_length=200)
    call_id: str = ""


async def _retell_transcript(call_id: str) -> list[dict[str, str]]:
    import httpx

    async with httpx.AsyncClient(timeout=20) as http:
        r = await http.get(f"{RETELL}/v2/get-call/{call_id}",
                           headers={"Authorization": f"Bearer {config.RETELL_API_KEY}"})
    if r.status_code >= 400:
        raise HTTPException(502, "Couldn't fetch the call from the voice service.")
    return transcript_of(r.json())


def transcript_of(call: dict[str, Any]) -> list[dict[str, Any]]:
    """Retell's transcript as turns, each with the second it started (from word timings)."""
    out = []
    for t in call.get("transcript_object") or []:
        if not _s(t.get("content")):
            continue
        words = t.get("words") or []
        start = words[0].get("start") if words and isinstance(words[0], dict) else None
        out.append({"role": "agent" if t.get("role") == "agent" else "user", "text": _s(t.get("content")),
                    "t": round(float(start), 1) if isinstance(start, (int, float)) else None})
    return out


@router.post("/agents/{agent_id}/score")
async def score(agent_id: str, body: ScoreBody, request: Request) -> dict[str, Any]:
    row = _row(agent_id, request)
    if body.call_id:
        test = next((t for t in row.get("tests") or [] if t.get("call_id") == body.call_id), None)
        if test is None:
            raise HTTPException(404, "No such test call for this agent.")
        if test.get("result"):
            return test["result"]
        transcript = await _retell_transcript(body.call_id)
    else:
        transcript = [m.model_dump() for m in body.transcript]
    if sum(1 for t in transcript if t["role"] == "user") < 2:
        raise HTTPException(422, "The participant needs to answer at least twice before this can be scored.")
    try:
        result = _evaluate(row, transcript)
    except LLMError as exc:
        raise _no_model() from exc
    except sim.EvaluationError as exc:
        raise HTTPException(422, {"message": "This test couldn't be evaluated: " + "; ".join(exc.problems),
                                  "problems": exc.problems}) from exc
    result["transcript"] = transcript
    if body.call_id:
        test["result"] = result
        store.save(row)
    return result


# --------------------------------------------------------------------------- #
#  Publish
# --------------------------------------------------------------------------- #
@router.post("/agents/{agent_id}/publish")
def publish(agent_id: str, request: Request) -> dict[str, Any]:
    row = _editable(agent_id, request)
    # Re-apply today's rules before freezing: an agent saved under older ones
    # (video, chat, "Interview") must not publish them.
    row["cfg"] = clean_cfg(row["cfg"])
    row["fields"] = clean_fields(row["fields"], row["cfg"])
    missing = [i["label"] for i in readiness(row) if not i["done"]]
    if missing:
        raise HTTPException(422, {"message": "Finish these before publishing.", "missing": missing})
    body = call_body(row)
    leaks = rx.leaked_cues(body, row)
    if leaks:
        raise HTTPException(422, {"message": "The instructions or questions repeat a skill description.", "leaked": leaks})
    version = int((row.get("published") or {}).get("version") or 0) + 1
    row["published"] = {"version": version, "at": time.time(), "fields": row["fields"],
                        "agent": row["agent"], "cfg": row["cfg"]}
    # The open link is NOT made here: it is off until someone switches it on
    # in the Invite window (see /open-link).
    return {**public(store.save(row)), "published": True}


# --------------------------------------------------------------------------- #
#  Participants: invites and results
# --------------------------------------------------------------------------- #
class InviteBody(BaseModel):
    name: str = Field(default="", max_length=120)
    email: str = Field(default="", max_length=200)
    message: str = Field(default="", max_length=1500)
    send_email: bool = False


def _code() -> str:
    import secrets
    import string

    digits = "".join(secrets.choice(string.digits) for _ in range(4))
    tail = "".join(secrets.choice("ABCDEFGHJKLMNPQRSTUVWXYZ") for _ in range(2))
    return f"RP-{digits}-{tail}"


def _unique_code() -> str:
    codes = {i["code"] for r in store.list_all() for i in r.get("invites") or []}
    code = _code()
    while code in codes:
        code = _code()
    return code


def _base(request: Request) -> str:
    return config.PUBLIC_URL or str(request.base_url).rstrip("/")


def invitation_email(row: dict[str, Any], name: str, code: str, link: str, note: str) -> tuple[str, str]:
    """(subject, text) of an invitation: what it is, how long, how to join."""
    a, cfg = row["published"]["agent"], row["published"]["cfg"]
    target, _ = rx.lengths(cfg)
    first = (name.split() or ["there"])[0]
    kind = "conversation"
    subject = f"Your invitation: {no_interview(a['title'])}"
    text = (
        f"Hi {first},\n\n"
        f"You're invited to a spoken {kind}: {no_interview(a['title'])}. It takes about {target} minutes and runs in your web browser "
        "on a computer or phone. You'll need a microphone and a quiet place.\n\n"
        + (note.strip() + "\n\n" if note.strip() else "")
        + f"Open your invitation: {link}\nYour access code: {code}\n\n"
        "Sign in with your name and this email address, choose a time, check your microphone, and start when you're ready.\n"
    )
    return subject, text


@router.post("/agents/{agent_id}/duplicate")
def duplicate(agent_id: str, request: Request) -> dict[str, Any]:
    """An editable draft copy of a role-play (published or not). It starts with no
    versions, invites, open link, tests or results: those stay with the original."""
    import copy

    src = _row(agent_id, request)
    agent = copy.deepcopy(src["agent"])
    agent["title"] = (agent["title"] + " (copy)")[:120]
    row = {
        "agent_id": store.new_id(agent["title"]), "org_id": src["org_id"], "brief": src.get("brief", ""),
        "mode": src.get("mode", "roleplay"), "fields": copy.deepcopy(src["fields"]), "agent": agent,
        "cfg": copy.deepcopy(src["cfg"]), "reviewed": False, "version": 0, "duplicated_from": agent_id,
    }
    return public(store.save(row))


@router.post("/agents/{agent_id}/invites")
def invite(agent_id: str, body: InviteBody, request: Request) -> dict[str, Any]:
    """One access code per participant, for the PUBLISHED version; optionally emailed."""
    row = _row(agent_id, request)
    if not row.get("published"):
        raise HTTPException(409, "Publish the agent before inviting anyone.")
    addr = _s(body.email).lower()
    if body.send_email and not re.fullmatch(r"[^@\s]+@[^@\s]+\.[^@\s]+", addr):
        raise HTTPException(422, "Enter a valid email address to send the invitation.")
    code = _unique_code()
    link = f"{_base(request)}/participant?code={code}"
    subject, text = invitation_email(row, _s(body.name), code, link, body.message)
    entry = {"code": code, "name": _s(body.name), "email": addr, "created_at": time.time(), "proctoring": _proctoring(row)}
    sent, error = False, ""
    if body.send_email:
        try:
            email.send(addr, subject, text)
            sent, entry["emailed_at"] = True, time.time()
        except email.EmailNotSent as exc:
            error = str(exc)
    row.setdefault("invites", []).append(entry)
    store.save(row)
    return {"code": code, "link": link, "version": row["published"]["version"],
            "sent": sent, "email_error": error, "subject": subject, "text": text}


@router.get("/agents/{agent_id}/invites")
def list_invites(agent_id: str, request: Request) -> dict[str, Any]:
    from services.data import agent_sessions

    row = _row(agent_id, request)
    by_code: dict[str, list[dict[str, Any]]] = {}
    for r in agent_sessions.for_agent(agent_id):
        by_code.setdefault(r.get("invite_code", ""), []).append(r)
    out = []
    for i in row.get("invites") or []:
        if i.get("shared"):
            continue
        rs = by_code.get(i["code"], [])
        status = ("Completed" if any(r["status"] == "complete" for r in rs)
                  else "Started" if rs else "Revoked" if i.get("revoked") else "Not started")
        out.append({"code": i["code"], "name": i.get("name", ""), "email": i.get("email", ""),
                    "created_at": i.get("created_at"), "emailed": bool(i.get("emailed_at")), "status": status,
                    "link": f"{_base(request)}/participant?code={i['code']}"})
    return {"invites": sorted(out, key=lambda x: x["created_at"] or 0, reverse=True), "open_link": _open_link(row)}


def _invite_values(row: dict[str, Any], name: str) -> dict[str, str]:
    pub = row["published"]
    target, _ = rx.lengths(pub["cfg"])
    return {"PARTICIPANT_NAME": (name.split() or ["there"])[0], "ROLE_PLAY": no_interview(pub["agent"]["title"]),
            "LANGUAGE": pub["cfg"].get("language") or rx.voice(pub["cfg"].get("voice") or "").language,
            "DURATION": str(target), "COMPANY_NAME": config.COMPANY_NAME}


@router.get("/agents/{agent_id}/invitation")
def invitation(agent_id: str, request: Request) -> dict[str, Any]:
    """Everything the Send Invitation window shows."""
    row = _row(agent_id, request)
    if not row.get("published"):
        raise HTTPException(409, "Publish the agent before inviting anyone.")
    purpose = _purpose(row)
    out = _open_link(row)
    if out["enabled"]:
        out["link"] = _base(request) + out["path"]
    return {"title": no_interview(row["published"]["agent"]["title"]), "purpose": purpose,
            "label": "AI Conversation",
            "template": row.get("invite_template") or tmpl.default_template(purpose),
            "placeholders": ["{" + p + "}" for p in tmpl.PLACEHOLDERS],
            "defaults": _proctoring(row), "open_link": out, "max_emails": 10,
            "credits": "Unlimited", "email_configured": email.configured(), "email_from": email.sender()}


class InvitationsBody(BaseModel):
    emails: list[str] = Field(min_length=1, max_length=10)
    template_html: str = Field(default="", max_length=20000)
    image_proctoring: bool = False
    safe_browser: bool = False


@router.post("/agents/{agent_id}/invitations")
def send_invitations(agent_id: str, body: InvitationsBody, request: Request) -> dict[str, Any]:
    """Up to ten people at once: each gets their own access code and, when email is
    set up, the invitation from the template with their link and code added."""
    row = _row(agent_id, request)
    if not row.get("published"):
        raise HTTPException(409, "Publish the agent before inviting anyone.")
    emails, seen = [], set()
    for e in body.emails:
        e = _s(e).lower()
        if e and e not in seen:
            seen.add(e)
            emails.append(e)
    bad = [e for e in emails if not re.fullmatch(r"[^@\s]+@[^@\s]+\.[^@\s]+", e)]
    if bad:
        raise HTTPException(422, "These aren't valid email addresses: " + ", ".join(bad))
    tpl = body.template_html.strip() or tmpl.default_template(_purpose(row))
    row["invite_template"] = tmpl.sanitize(tpl)
    proctoring = _proctoring(row, body.image_proctoring, body.safe_browser)
    subject = f"Your invitation: {no_interview(row['published']['agent']['title'])}"
    results = []
    for e in emails:
        code = _unique_code()
        link = f"{_base(request)}/participant?code={code}"
        name = e.split("@")[0].replace(".", " ").title()
        html_ = (tmpl.fill(tpl, _invite_values(row, name))
                 + f'<p>Your link: <a href="{link}">{link}</a><br>Your access code: <b>{code}</b></p>')
        entry = {"code": code, "name": "", "email": e, "created_at": time.time(), "proctoring": proctoring}
        sent, err = False, ""
        if email.configured():
            try:
                email.send(e, subject, tmpl.to_text(html_), html_)
                sent, entry["emailed_at"] = True, time.time()
            except email.EmailNotSent as exc:
                err = str(exc)
        row.setdefault("invites", []).append(entry)
        results.append({"email": e, "code": code, "link": link, "sent": sent, "error": err,
                        "text": tmpl.to_text(html_)})
    store.save(row)
    return {"subject": subject, "results": results, "email_configured": email.configured()}


class OpenLinkBody(BaseModel):
    enabled: bool
    image_proctoring: bool | None = None
    safe_browser: bool | None = None


@router.post("/agents/{agent_id}/open-link")
def open_link(agent_id: str, body: OpenLinkBody, request: Request) -> dict[str, Any]:
    """Switch the agent's open link on or off. Off means its code stops working at once."""
    row = _row(agent_id, request)
    if body.enabled and not row.get("published"):
        raise HTTPException(409, "Publish the agent before opening a link to it.")
    inv = next((i for i in row.get("invites") or [] if i.get("shared")), None)
    if body.enabled:
        if inv is None:
            inv = {"code": _unique_code(), "shared": True, "created_at": time.time()}
            row.setdefault("invites", []).append(inv)
        else:
            inv.pop("revoked", None)
    elif inv is not None:
        inv["revoked"] = True
    if inv is not None:
        inv["proctoring"] = _proctoring(row, body.image_proctoring, body.safe_browser, inv.get("proctoring"))
    store.save(row)
    out = _open_link(row)
    if out["enabled"]:
        out["link"] = _base(request) + out["path"]
    return out


@router.get("/agents/{agent_id}/sessions")
def participant_sessions(agent_id: str, request: Request) -> dict[str, Any]:
    from services.data import agent_sessions

    _row(agent_id, request)
    return {"sessions": [
        {"session_id": r["session_id"], "name": r.get("name", ""), "email": r.get("email", ""),
         "status": r["status"], "version": r["version"], "early": bool(r.get("early")),
         "elapsed_sec": r.get("elapsed_sec", 0), "ended_at": r.get("ended_at"),
         "attempt": int(r.get("attempt") or 1),
         "evaluation": r.get("evaluation"), "evaluation_error": r.get("evaluation_error"),
         "result": r.get("result"), "result_error": r.get("result_error"),
         "feedback": r.get("feedback"), "transcript": r.get("transcript") or [],
         "booking": r.get("booking")}
        for r in agent_sessions.for_agent(agent_id)
    ]}


# --------------------------------------------------------------------------- #
#  Reports: one grid per role-play, a report per attempt, admin review
# --------------------------------------------------------------------------- #
#: The admin's own recommendation, by purpose. The AI's recommendation sits next
#: to it and is never overwritten: the decision is a person's.
DECISIONS = {
    sim.GENERAL: ("Recommended", "Needs review", "Not recommended"),
    "Hiring": ("Advance", "Hold", "Reject"),
    "HR": ("No action", "Follow up", "Escalate"),
    "L&D": ("Ready", "Needs practice", "Coaching recommended"),
}


def _purpose(row: dict[str, Any]) -> str:
    """Every role-play is General: Purpose was removed until there's data to define it."""
    return sim.GENERAL


#: Icon level for a recommendation: positive (check), caution (!), negative (x).
_LEVEL = {
    "Recommended": "positive", "Needs review": "caution", "Not recommended": "negative",
    "Advance": "positive", "Hold": "caution", "Reject": "negative",
    "No action": "positive", "Follow up": "caution", "Escalate": "negative",
    "Ready": "positive", "Needs practice": "caution", "Coaching recommended": "negative",
}


def _proctoring_label(p: dict[str, Any] | None) -> str:
    p = p or {}
    on = [n for k, n in (("image_proctoring", "Image"), ("safe_browser", "Safe browser")) if p.get(k)]
    return " + ".join(on) or "Off"


def _ai_level(rec: str) -> str:
    r = (rec or "").lower()
    if r.startswith(("recommended", "proceed", "advanced", "proficient")):
        return "positive"
    if r.startswith(("not recommended", "not suitable", "foundational")):
        return "negative"
    return "caution" if r else ""


def _report_row(r: dict[str, Any], skills: list[str]) -> dict[str, Any]:
    ev = r.get("evaluation") or {}
    by = {x["name"]: x for x in ev.get("skills") or []}
    started, ended = r.get("created_at"), r.get("ended_at")
    recent = (time.time() - (ended or 0)) < 600
    status = ("In Progress" if r["status"] != "complete" else "Terminated" if r.get("early") else "Completed")
    evaluation = ("Evaluated" if ev else "" if r["status"] != "complete"
                  else "Evaluating" if recent and not r.get("evaluation_error") else "Not evaluated")
    review = r.get("review") or {}
    rec, src = (review["decision"], "admin") if review.get("decision") else (ev.get("recommendation", ""), "ai" if ev else "")
    return {
        "session_id": r["session_id"], "name": r.get("name", ""), "email": r.get("email", ""),
        "attempt": int(r.get("attempt") or 1), "version": r.get("version"), "status": status, "evaluation": evaluation,
        "started_at": started, "ended_at": ended, "date": ended or started,
        "duration_sec": int(r.get("elapsed_sec") or 0) or (int(ended - started) if ended and started else 0),
        "early": bool(r.get("early")),
        "overall": ev.get("overall"), "band": ev.get("band", ""), "ai_recommendation": ev.get("recommendation", ""),
        "recommendation": rec, "recommendation_source": src,
        "recommendation_level": _LEVEL.get(rec, "") if src == "admin" else _ai_level(rec),
        "proctoring": _proctoring_label(r.get("proctoring")),
        "coverage": ev.get("weight_coverage"),
        "skills": {n: (by[n]["score"] if n in by else None) for n in skills},
        "review": review,
        "problems": (r.get("evaluation_error") or {}).get("problems", []),
    }


def _pending_rows(row: dict[str, Any], started_codes: set[str]) -> list[dict[str, Any]]:
    """People invited by email who haven't started: they belong in the grid too."""
    out = []
    for i in row.get("invites") or []:
        if i.get("shared") or i.get("revoked") or i["code"] in started_codes:
            continue
        out.append({"session_id": "", "invite_code": i["code"], "name": i.get("name", ""), "email": i.get("email", ""),
                    "attempt": 0, "status": "Pending", "evaluation": "", "date": i.get("created_at"),
                    "started_at": None, "ended_at": None, "duration_sec": 0, "early": False, "overall": None, "band": "",
                    "ai_recommendation": "", "recommendation": "", "recommendation_source": "", "recommendation_level": "",
                    "proctoring": _proctoring_label(i.get("proctoring")),
                    "coverage": None, "skills": {}, "review": {}, "problems": []})
    return out


def _agent_report(row: dict[str, Any]) -> dict[str, Any]:
    from services.data import agent_sessions

    a = (row.get("published") or {}).get("agent") or row["agent"]
    purpose = _purpose(row)
    skills = [{"name": r["name"], "weight": r["weight"]} for r in a["rubric"]]
    names = [x["name"] for x in skills]
    sessions_ = agent_sessions.for_agent(row["agent_id"])
    rows = [_report_row(r, names) for r in sessions_]
    rows += _pending_rows(row, {r.get("invite_code", "") for r in sessions_})
    for x in rows:
        x["agent_id"], x["agent_title"] = row["agent_id"], no_interview(a["title"])
    done = [x for x in rows if x["overall"] is not None]
    avg = lambda vals: round(sum(vals) / len(vals), 1) if vals else None  # noqa: E731
    return {
        "agent": {"agent_id": row["agent_id"], "title": no_interview(a["title"]), "type_label": no_interview(a.get("type_label", "")), "purpose": purpose,
                  "role": ((row.get("published") or {}).get("fields") or row["fields"]).get("role", ""),
                  "published_version": int((row.get("published") or {}).get("version") or 0), "skills": skills},
        "decisions": DECISIONS.get(purpose, DECISIONS["L&D"]),
        "stats": {
            "invited": sum(1 for i in row.get("invites") or [] if not i.get("shared")),
            "attempts": sum(1 for x in rows if x["status"] != "Pending"),
            "people": len({x["email"] for x in rows if x["email"]}),
            "completed": sum(1 for x in rows if x["status"] in ("Completed", "Terminated")),
            "evaluated": len(done), "average": avg([x["overall"] for x in done]),
            "skills": {n: avg([x["skills"][n] for x in done if x["skills"].get(n) is not None]) for n in names},
            "decisions": {d: sum(1 for x in rows if x["review"].get("decision") == d) for d in DECISIONS.get(purpose, ())},
            "reviewed": sum(1 for x in rows if x["review"].get("decision")),
        },
        "rows": sorted(rows, key=lambda x: x["date"] or 0, reverse=True),
    }


@router.get("/agents/{agent_id}/report")
def report(agent_id: str, request: Request) -> dict[str, Any]:
    """Every attempt and pending invitation for this role-play, with totals."""
    return _agent_report(_row(agent_id, request))


@router.get("/reports")
def all_reports(request: Request) -> dict[str, Any]:
    """Every attempt across this organization's published role-plays, newest first."""
    org = _org(request)
    rows, agents_ = [], []
    for r in store.list_all():
        if r.get("org_id") != org or not r.get("published"):
            continue
        rep = _agent_report(r)
        rows += rep["rows"]
        agents_.append({"agent_id": r["agent_id"], "title": rep["agent"]["title"], "purpose": rep["agent"]["purpose"]})
    rows.sort(key=lambda x: x["date"] or 0, reverse=True)
    return {"rows": rows, "agents": sorted(agents_, key=lambda a: a["title"].lower())}


def _filtered(rows: list[dict[str, Any]], q: str, status: str, rec: str, agent: str) -> list[dict[str, Any]]:
    q = (q or "").strip().lower()
    return [x for x in rows if (not q or q in (x["name"] + " " + x["email"]).lower())
            and (not status or x["status"] == status) and (not rec or x["recommendation_level"] == rec)
            and (not agent or x["agent_id"] == agent)]


def _xlsx(rows: list[dict[str, Any]], skills: list[str], title: str) -> Response:
    import io
    from datetime import datetime

    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "Reports"
    head = (["Participant", "Email", "Role-play", "Attempt", "Date", "Status", "Duration (min)", "Score", "Band",
             "AI recommendation", "Decision", "Decision by", "Evaluation notes", "Proctoring"] + skills)
    ws.append(head)
    for c in ws[1]:
        c.font, c.fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="7F56D9")
    for x in rows:
        when = datetime.fromtimestamp(x["date"]).strftime("%d %b %Y %H:%M") if x["date"] else ""
        ws.append([x["name"], x["email"], x["agent_title"], x["attempt"] or "", when, x["status"],
                   round(x["duration_sec"] / 60, 1) if x["duration_sec"] else "", x["overall"], x["band"],
                   x["ai_recommendation"], x["review"].get("decision", ""), x["review"].get("by", ""),
                   x["review"].get("notes", ""), x["proctoring"]] + [x["skills"].get(n) for n in skills])
    for col, width in zip("ABCDEFGHIJKLMN", (22, 30, 30, 8, 18, 12, 14, 8, 12, 36, 14, 26, 50, 11)):
        ws.column_dimensions[col].width = width
    ws.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    name = re.sub(r"[^\w-]+", "_", title)[:60] or "reports"
    return Response(buf.getvalue(), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": f'attachment; filename="{name}.xlsx"'})


@router.get("/agents/{agent_id}/report.xlsx")
def report_xlsx(agent_id: str, request: Request, q: str = "", status: str = "", rec: str = "") -> Response:
    rep = _agent_report(_row(agent_id, request))
    return _xlsx(_filtered(rep["rows"], q, status, rec, ""), [s["name"] for s in rep["agent"]["skills"]], rep["agent"]["title"] + " reports")


@router.get("/reports.xlsx")
def all_reports_xlsx(request: Request, q: str = "", status: str = "", rec: str = "", agent: str = "") -> Response:
    data = all_reports(request)
    return _xlsx(_filtered(data["rows"], q, status, rec, agent), [], "Role-play reports")


_REGIONS = {"US": "United States", "GB": "United Kingdom", "IN": "India", "AU": "Australia", "CA": "Canada",
            "FR": "France", "DE": "Germany", "ES": "Spain", "MX": "Mexico", "AE": "United Arab Emirates"}


def _language_label(cfg: dict[str, Any]) -> str:
    """'English (United States)' from the voice's locale."""
    v = rx.voice(cfg.get("voice") or "")
    region = (v.locale.split("-") + [""])[1].upper()
    return f"{v.language} ({_REGIONS[region]})" if region in _REGIONS else v.language


def _session_of(agent_id: str, session_id: str, request: Request) -> tuple[dict[str, Any], dict[str, Any]]:
    from services.data import agent_sessions

    row = _row(agent_id, request)
    s = agent_sessions.load(session_id)
    if s is None or s.get("agent_id") != agent_id:
        raise HTTPException(404, "No such attempt.")
    return row, s


# "attempt_id", not "session_id": a session_id path parameter is checked against
# the interview product's sessions; a role-play attempt is checked by _session_of.
@router.get("/agents/{agent_id}/attempts/{attempt_id}/report")
def session_report(agent_id: str, attempt_id: str, request: Request) -> dict[str, Any]:
    row, s = _session_of(agent_id, attempt_id, request)
    a = (row.get("published") or {}).get("agent") or row["agent"]
    cfg = (s.get("snapshot") or {}).get("cfg") or (row.get("published") or {}).get("cfg") or row["cfg"]
    inv = next((i for i in row.get("invites") or [] if i.get("code") == s.get("invite_code")), {})
    return {"session_id": attempt_id, "agent_id": agent_id, "agent_title": no_interview(a["title"]),
            "language": _language_label(cfg),
            "invitation_type": "Open Link Invite" if inv.get("shared") else "Email Invite" if inv else "",
            "date": s.get("ended_at") or s.get("created_at"), "duration_sec": int(s.get("elapsed_sec") or 0),
            "persona": (a.get("persona") or {}).get("name", ""), "purpose": _purpose(row),
            "decisions": DECISIONS.get(_purpose(row), DECISIONS["L&D"]), "name": s.get("name", ""), "email": s.get("email", ""),
            "attempt": int(s.get("attempt") or 1), "status": s["status"], "evaluation": s.get("evaluation"),
            "evaluation_error": s.get("evaluation_error"), "transcript": s.get("transcript") or [],
            "review": s.get("review") or {}, "feedback": s.get("feedback")}


@router.post("/agents/{agent_id}/attempts/{attempt_id}/evaluate")
async def evaluate_now(agent_id: str, attempt_id: str, request: Request) -> dict[str, Any]:
    """Run (or re-run) the evaluation of one finished attempt, e.g. one taken before
    the evaluation engine existed, or one whose evaluation failed."""
    from services.api import participant as part
    from services.data import agent_sessions

    _, s = _session_of(agent_id, attempt_id, request)
    if s["status"] != "complete":
        raise HTTPException(409, "This attempt hasn't finished yet.")
    if s.get("calls") and config.RETELL_API_KEY:
        s["transcript"] = await part._transcripts(s["calls"]) or s.get("transcript") or []
    s.pop("evaluation", None)
    part.evaluate_session(s)
    agent_sessions.save(s)
    if not s.get("evaluation"):
        raise HTTPException(422, "Couldn't evaluate: " + "; ".join((s.get("evaluation_error") or {}).get("problems") or ["unknown"]))
    return {"ok": True}


@router.get("/agents/{agent_id}/attempts/{attempt_id}/audio")
async def attempt_audio(agent_id: str, attempt_id: str, request: Request) -> dict[str, Any]:
    """The recording of each call in this attempt (a reconnect is a second call),
    fetched fresh from the voice service so links never go stale in a report."""
    import httpx

    _, s = _session_of(agent_id, attempt_id, request)
    out = []
    if config.RETELL_API_KEY:
        async with httpx.AsyncClient(timeout=20) as http:
            for n, cid in enumerate(s.get("calls") or [], 1):
                try:
                    r = await http.get(f"{RETELL}/v2/get-call/{cid}", headers={"Authorization": f"Bearer {config.RETELL_API_KEY}"})
                except Exception:  # noqa: BLE001 — one missing recording shouldn't hide the others
                    continue
                if r.status_code < 400 and (r.json().get("recording_url") or ""):
                    c = r.json()
                    out.append({"part": n, "url": c["recording_url"],
                                "duration_sec": round(((c.get("end_timestamp") or 0) - (c.get("start_timestamp") or 0)) / 1000)})
    return {"recordings": out}


def exhibit_dir():
    d = config.DATA_DIR / "exhibits"
    d.mkdir(parents=True, exist_ok=True)
    return d


_MAGIC = ((b"\x89PNG\r\n\x1a\n", "png"), (b"\xff\xd8\xff", "jpg"), (b"RIFF", "webp"))


class ExhibitImageBody(BaseModel):
    data_base64: str = Field(min_length=10, max_length=7_000_000)   # ~5 MB of image


@router.post("/agents/{agent_id}/exhibit-images")
def upload_exhibit_image(agent_id: str, body: ExhibitImageBody, request: Request) -> dict[str, Any]:
    """An image for an exhibit: PNG, JPEG or WebP, up to 5 MB, checked by its bytes, not its name."""
    import base64
    import secrets as _secrets

    _editable(agent_id, request)
    try:
        raw = base64.b64decode(body.data_base64.split(",", 1)[-1], validate=True)
    except ValueError as exc:
        raise HTTPException(422, "That file couldn't be read.") from exc
    ext = next((e for m, e in _MAGIC if raw.startswith(m) and (e != "webp" or raw[8:12] == b"WEBP")), None)
    if ext is None or len(raw) > 5 * 1024 * 1024:
        raise HTTPException(422, "Use a PNG, JPEG or WebP image up to 5 MB.")
    name = f"{_secrets.token_hex(12)}.{ext}"
    (exhibit_dir() / name).write_bytes(raw)
    return {"file": name}


def _exhibit_files(snapshot_agent: dict[str, Any]) -> set[str]:
    return {e.get("file") for e in snapshot_agent.get("exhibits") or [] if e.get("file")}


def exhibit_response(name: str):
    from fastapi.responses import FileResponse

    path = exhibit_dir() / name
    if not _FILE_ID.match(name) or not path.exists():
        raise HTTPException(404, "No such image.")
    media = {"png": "image/png", "jpg": "image/jpeg", "webp": "image/webp"}[name.rsplit(".", 1)[1]]
    return FileResponse(path, media_type=media, headers={"Cache-Control": "private, max-age=3600",
                                                         "X-Content-Type-Options": "nosniff"})


@router.get("/agents/{agent_id}/exhibit-images/{name}")
def get_exhibit_image(agent_id: str, name: str, request: Request):
    """Only an image this agent (its draft or published version) actually uses."""
    row = _row(agent_id, request)
    if name not in _exhibit_files(row["agent"]) | _exhibit_files((row.get("published") or {}).get("agent") or {}):
        raise HTTPException(404, "No such image.")
    return exhibit_response(name)


class ReviewBody(BaseModel):
    decision: str = Field(default="", max_length=40)
    notes: str = Field(default="", max_length=5000)


@router.put("/agents/{agent_id}/attempts/{attempt_id}/review")
def review(agent_id: str, attempt_id: str, body: ReviewBody, request: Request) -> dict[str, Any]:
    """The admin's evaluation notes and recommendation for one attempt."""
    from services.data import agent_sessions
    from services.security import principal as security

    row, s = _session_of(agent_id, attempt_id, request)
    allowed = DECISIONS.get(_purpose(row), DECISIONS["L&D"])
    if body.decision and body.decision not in allowed:
        raise HTTPException(422, "Choose one of: " + ", ".join(allowed) + ".")
    who = security.optional_principal(request)
    s["review"] = {"decision": body.decision, "notes": body.notes.strip(),
                   "by": getattr(who, "email", "") or "Admin", "at": time.time()}
    agent_sessions.save(s)
    return s["review"]


# --------------------------------------------------------------------------- #
#  Retell webhook
# --------------------------------------------------------------------------- #
def verify_signature(raw: bytes, header: str, api_key: str, now_ms: int | None = None) -> bool:
    """Retell's `x-retell-signature`: `v=<ms>,d=<hex>` where hex is
    HMAC-SHA256(key=api key, msg=raw body + timestamp). Five-minute window."""
    try:
        parts = dict(p.split("=", 1) for p in header.split(","))
        ts, digest = parts["v"], parts["d"]
        age = abs((now_ms if now_ms is not None else int(time.time() * 1000)) - int(ts))
    except (KeyError, ValueError):
        return False
    if age > 5 * 60 * 1000:
        return False
    expected = hmac.new(api_key.encode(), raw + ts.encode(), hashlib.sha256).hexdigest()
    return hmac.compare_digest(expected, digest)


def _evaluate(row: dict[str, Any], transcript: list[dict[str, Any]]) -> dict[str, Any]:
    """A builder test, scored by the same engine and rules as a real session."""
    from services.assessment.agent_builder_flow import FLOW

    return sim.evaluate({**row, "cfg": clean_cfg(row["cfg"])}, transcript, complete=ab._complete, flow=FLOW)


def _score_in_background(agent_id: str, call_id: str, transcript: list[dict[str, str]]) -> None:
    row = store.load(agent_id)
    if row is None:
        return
    test = next((t for t in row.get("tests") or [] if t.get("call_id") == call_id), None)
    if test is None or test.get("result"):
        return
    test["transcript"] = transcript
    if sum(1 for t in transcript if t["role"] == "user") >= 2:
        try:
            test["result"] = {**_evaluate(row, transcript), "transcript": transcript}
        except Exception:  # noqa: BLE001 — scoring can be retried from the page
            pass
    store.save(row)


@webhook_router.post("/api/agent-builder/retell-webhook")
async def retell_webhook(request: Request, background: BackgroundTasks) -> dict[str, Any]:
    raw = await request.body()
    if not config.RETELL_API_KEY or not verify_signature(
        raw, request.headers.get("x-retell-signature", ""), config.RETELL_API_KEY
    ):
        raise HTTPException(401, "Bad signature.")
    try:
        event = json.loads(raw)
    except json.JSONDecodeError as exc:
        raise HTTPException(400, "Not JSON.") from exc
    call = event.get("call") or {}
    meta = call.get("metadata") or {}
    if event.get("event") not in ("call_ended", "call_analyzed") or meta.get("source") != "tara-agent-builder":
        return {"ok": True, "ignored": True}
    call_id = _s(call.get("call_id"))
    # The agent comes from OUR record of the call, never from metadata alone:
    # metadata is what we sent, but only the binding proves we placed the call.
    row = store.by_call(call_id)
    if row is None:
        return {"ok": True, "ignored": True}
    background.add_task(_score_in_background, row["agent_id"], call_id, transcript_of(call))
    return {"ok": True}
