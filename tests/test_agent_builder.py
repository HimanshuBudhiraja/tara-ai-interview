"""The Agent Builder and its one Retell agent.

What must hold, whatever the model writes:
  * every call sends all 18 variables, as strings, and the prompt is left with no `{{`;
  * voice comes only from Persona → Voice, and length only from Persona → Follow-up depth;
  * the rubric never reaches Retell, and a payload that would carry it is refused;
  * nothing in the code can create a Retell agent: there is one;
  * the webhook acts only on a correctly signed body;
  * one organization never sees another's agents.
"""
from __future__ import annotations

import copy
import hashlib
import hmac
import json
import re
import time
from pathlib import Path

import pytest
from fastapi.testclient import TestClient

from services import config
from services.ai.workloads import agent_builder as ab
from services.api import agent_builder as api
from services.api.app import app
from services.assessment import agent_builder_retell as rx
from services.data import built_agents as store
from tests.conftest import sign_in

BASE = "/api/recruiter/agent-builder"


def _row(**cfg) -> dict:
    p = ab.normalise_plan(ab._offline_plan("A technical interview for a senior AI Engineer.", "interview"))
    c = ab.normalise_content(ab._offline_content(p))
    row = api._new_row(p, c, "brief", "interview", "org")
    row["cfg"] = api.clean_cfg({**row["cfg"], **cfg})
    return row


@pytest.fixture()
def client(data_dir, tenant, monkeypatch):
    from services.ai import brain

    def _no_model(*_a, **_k):
        raise brain.LLMError("offline")

    monkeypatch.setattr(brain.RuntimeBrain, "complete_json", _no_model)
    # The Agent Builder calls the gateway on its own workloads; keep every test offline.
    monkeypatch.setattr(ab, "_complete", _no_model)
    c = TestClient(app)
    sign_in(c, tenant)
    return c


def _draft(c, brief="A 20-minute technical interview for a senior AI Engineer.", mode="interview"):
    r = c.post(f"{BASE}/drafts", json={"brief": brief, "mode": mode})
    assert r.status_code == 200, r.text
    events = {}
    for chunk in r.text.strip().split("\n\n"):
        ev = re.search(r"^event: (.*)$", chunk, re.M).group(1)
        events[ev] = json.loads(re.search(r"^data: (.*)$", chunk, re.M).group(1))
    assert "error" not in events, events.get("error")
    return events


# --------------------------------------------------------------------------- #
#  The variable contract and the prompt
# --------------------------------------------------------------------------- #
def test_the_flow_uses_only_contract_and_system_variables():
    import json as _json

    from services.assessment import agent_builder_flow as flow

    used = rx.prompt_variables() | rx.prompt_variables(_json.dumps(flow.NODES))
    assert used <= set(rx.VARIABLES) | rx.SYSTEM_VARIABLES
    # Sent but read by no node: max_minutes sets the call cap, language goes as
    # agent_override.agent.language.
    assert set(rx.VARIABLES) - used <= {"max_minutes", "language"}


def test_the_flow_closes_only_after_the_wrap_up():
    from services.assessment import agent_builder_flow as flow

    nodes = {n["id"]: n for n in flow.NODES}
    assert list(nodes) == ["opening", "resume", "main", "followup", "wrapup", "stop", "closing", "end_call"]
    assert flow.FLOW["start_node_id"] == "opening"
    out = {k: {e["destination_node_id"] for e in n.get("edges", [])} | (
        {n["skip_response_edge"]["destination_node_id"]} if n.get("skip_response_edge") else set()) for k, n in nodes.items()}
    assert out["opening"] == {"resume", "main", "stop"} and out["main"] == {"followup", "wrapup", "stop"}
    assert out["followup"] == {"main", "wrapup", "stop"} and out["stop"] == {"closing"} and out["closing"] == {"end_call"}
    # Only the wrap-up (participant has nothing more, or out of time) and a stop reach the closing line.
    assert out["wrapup"] == {"closing", "stop"}
    assert {k for k, v in out.items() if "closing" in v} == {"wrapup", "stop"}
    assert "Tara decides" not in json.dumps(flow.FLOW)   # the mode reaches the agent as "<persona> decides"
    assert nodes["closing"]["instruction"] == {"type": "static_text", "text": "{{closing_line}}"}
    assert rx.general_prompt() == flow.GLOBAL_PROMPT.strip()      # the repo copy is the flow's prompt


@pytest.mark.parametrize("speaker", rx.SPEAKERS)
def test_every_call_fills_every_placeholder(speaker):
    row = _row(speaker=speaker)
    variables = rx.dynamic_variables(row)
    assert set(variables) == set(rx.VARIABLES)
    assert all(isinstance(v, str) for v in variables.values())
    assert "{{" not in rx.render(variables)


@pytest.mark.parametrize("engine,key", [("llm", "retell_llm"), ("flow", "conversation_flow")])
def test_the_persona_always_opens(engine, key):
    # "Who speaks first" was removed from the builder: whatever an old config says, the persona opens.
    body = rx.web_call_body({**_row(), "cfg": api.clean_cfg({**_row()["cfg"], "speaker": "Participant opens"})}, "agent_x", engine=engine)
    assert body["agent_override"][key]["start_speaker"] == "agent"
    assert body["retell_llm_dynamic_variables"]["opening_line"]
    if engine == "flow":
        # A flow takes no begin message: the opening travels as {{opening_line}}.
        assert set(body["agent_override"][key]) == {"start_speaker"}


# --------------------------------------------------------------------------- #
#  Voice and length come from Persona and nothing else
# --------------------------------------------------------------------------- #
@pytest.mark.parametrize("key", rx.VOICE_KEYS)
def test_the_voice_is_the_persona_voice_choice(key):
    body = rx.web_call_body(_row(voice=key), "agent_x")
    v = rx.voice(key)
    assert body["agent_override"]["agent"]["voice_id"] == v.voice_id
    assert body["agent_override"]["agent"]["language"] == v.locale
    assert body["retell_llm_dynamic_variables"]["language"] == v.language


@pytest.mark.parametrize("depth,minutes", [("Light", 10), ("Probing", 20), ("Deep dive", 30)])
@pytest.mark.parametrize("ending,cap", [("Tara decides", 5), ("Hard time limit", 0), ("No end time", None)])
def test_the_length_is_the_follow_up_depth(depth, minutes, ending, cap):
    body = rx.web_call_body(_row(depth=depth, ending=ending), "agent_x")
    expect_cap = 60 if cap is None else minutes + cap
    assert body["retell_llm_dynamic_variables"]["target_minutes"] == str(minutes)
    assert body["retell_llm_dynamic_variables"]["max_minutes"] == str(expect_cap)
    assert body["agent_override"]["agent"]["max_call_duration_ms"] == (expect_cap + 2) * 60_000


def test_difficulty_and_question_count_change_neither_voice_nor_length():
    base = rx.web_call_body(_row(voice="willa", depth="Probing"), "agent_x")
    row = _row(voice="willa", depth="Probing", tone="Tough")
    row["agent"]["questions"] *= 3
    other = rx.web_call_body(row, "agent_x")
    assert other["agent_override"]["agent"] == base["agent_override"]["agent"]  # tuning included
    assert other["retell_llm_dynamic_variables"]["target_minutes"] == "20"


def test_the_language_always_follows_the_voice():
    cfg = api.clean_cfg({"voice": "emma", "language": "German"})
    assert cfg["language"] == "French"


def test_every_voice_has_a_real_retell_id():
    for v in rx.VOICES:
        assert v.default_id and "RETELL_VOICE_ID" not in v.default_id, v.key
        assert "-" in v.default_id  # provider-Name, as Retell names them


# --------------------------------------------------------------------------- #
#  The rubric never crosses
# --------------------------------------------------------------------------- #
def test_a_clean_agent_sends_no_rubric():
    row = _row()
    body = rx.web_call_body(row, "agent_x")
    assert rx.leaked_cues(body, row) == []
    blob = json.dumps(body)
    for r in row["agent"]["rubric"]:
        assert r["anchor"] not in blob
    assert row["agent"]["description"] not in blob


def test_instructions_that_restate_the_rubric_are_caught():
    row = _row()
    row["agent"]["instructions"] += " " + row["agent"]["rubric"][0]["anchor"]
    assert rx.leaked_cues(rx.web_call_body(row, "agent_x"), row)


def test_nothing_in_the_code_creates_a_retell_agent():
    root = Path(config.ROOT_DIR)
    for path in list((root / "services").rglob("*.py")) + list((root / "tools").glob("retell*.py")):
        text = path.read_text()
        for verb in ("create-agent", "create-retell-llm", "create_agent("):
            assert verb not in text, f"{path} can create a Retell agent"


# --------------------------------------------------------------------------- #
#  The API
# --------------------------------------------------------------------------- #
def test_draft_streams_plan_then_content_then_the_saved_agent(client):
    ev = _draft(client)
    assert list(ev) == ["plan", "content", "done"]
    done = ev["done"]
    assert done["agent"]["questions"] and done["agent"]["rubric"]
    assert sum(r["weight"] for r in done["agent"]["rubric"]) == 100
    assert done["leaked_cues"] == []
    assert done["call_config"]["agent_override"]["agent"]["voice_id"] == rx.voice(done["cfg"]["voice"]).voice_id
    assert client.get(f"{BASE}/agents/{done['agent_id']}").status_code == 200


def test_save_publish_cycle(client):
    row = _draft(client)["done"]
    aid = row["agent_id"]
    r = client.post(f"{BASE}/agents/{aid}/publish")
    assert r.status_code == 422 and "Skills" in r.json()["detail"]["missing"]

    row["cfg"]["depth"] = "Deep dive"
    row["cfg"]["voice"] = "anthony"
    saved = client.put(f"{BASE}/agents/{aid}", json={**{k: row[k] for k in ("fields", "agent", "cfg")}, "reviewed": True})
    assert saved.status_code == 200, saved.text
    s = saved.json()
    assert s["length"] == {"target_minutes": 30, "cap_minutes": 35}
    assert s["fields"]["length"] == "30 min"
    assert s["call_config"]["agent_override"]["agent"]["voice_id"] == "11labs-Anthony"

    pub = client.post(f"{BASE}/agents/{aid}/publish")
    assert pub.status_code == 200, pub.text
    assert pub.json()["published_version"] == 1


def test_revise_needs_a_model(client):
    aid = _draft(client)["done"]["agent_id"]
    r = client.post(f"{BASE}/agents/{aid}/revise", json={"instruction": "make it tougher"})
    assert r.status_code == 503


def test_revise_applies_the_change_and_unreviews_the_rubric(client, monkeypatch):
    row = _draft(client)["done"]
    aid = row["agent_id"]
    client.put(f"{BASE}/agents/{aid}", json={**{k: row[k] for k in ("fields", "agent", "cfg")}, "reviewed": True})
    changed = copy.deepcopy(row["agent"])
    changed["persona"]["style"] = "blunt, skeptical"
    from services.ai import brain
    monkeypatch.setattr(ab, "_complete", lambda *a, **k: {
        "scenario": row["fields"], "agent": changed,
        "settings": {"difficulty": "Tough", "followUpDepth": "Light", "voice": "carola"},
        "summary": "Made it tougher and shorter.",
    })
    r = client.post(f"{BASE}/agents/{aid}/revise", json={"instruction": "tougher, shorter"})
    assert r.status_code == 200, r.text
    out = r.json()
    assert out["cfg"]["tone"] == "Tough" and out["cfg"]["depth"] == "Light"
    assert out["length"]["target_minutes"] == 10
    assert out["cfg"]["voice"] == "carola" and out["cfg"]["language"] == "German"
    assert out["reviewed"] is False
    assert out["summary"] == "Made it tougher and shorter."


def test_a_test_call_goes_to_the_one_agent_with_the_persona_overrides(client, monkeypatch):
    row = _draft(client)["done"]
    aid = row["agent_id"]
    monkeypatch.setattr(config, "RETELL_API_KEY", "key_test")
    monkeypatch.setattr(config, "RETELL_AGENT_BUILDER_AGENT_ID", "agent_one")
    sent = {}

    class FakeResponse:
        status_code = 201

        def json(self):
            return {"call_id": "call_123", "access_token": "tok"}

    class FakeClient:
        def __init__(self, *a, **k):
            pass

        async def __aenter__(self):
            return self

        async def __aexit__(self, *a):
            return False

        async def post(self, url, json=None, headers=None):
            sent.update(url=url, body=json)
            return FakeResponse()

    import httpx
    monkeypatch.setattr(httpx, "AsyncClient", FakeClient)
    r = client.post(f"{BASE}/agents/{aid}/test-call")
    assert r.status_code == 200, r.text
    assert r.json() == {"access_token": "tok", "call_id": "call_123", "max_minutes": rx.TEST_CALL_MINUTES}
    # A test is the real agent, kept short, and the persona is told so it can wrap up.
    assert sent["body"]["agent_override"]["agent"]["max_call_duration_ms"] == rx.TEST_CALL_MINUTES * 60_000
    assert sent["body"]["retell_llm_dynamic_variables"]["max_minutes"] == str(rx.TEST_CALL_MINUTES)
    assert sent["url"].endswith("/v3/create-web-call")      # v2 is retired on 2026-10-18
    assert sent["body"]["agent_id"] == "agent_one"
    assert set(sent["body"]["retell_llm_dynamic_variables"]) == set(rx.VARIABLES)
    assert store.by_call("call_123")["agent_id"] == aid


def test_a_test_call_that_would_leak_the_rubric_is_refused(client, monkeypatch):
    row = _draft(client)["done"]
    row["agent"]["instructions"] += " " + row["agent"]["rubric"][0]["anchor"]
    client.put(f"{BASE}/agents/{row['agent_id']}", json={**{k: row[k] for k in ("fields", "agent", "cfg")}})
    monkeypatch.setattr(config, "RETELL_API_KEY", "key_test")
    monkeypatch.setattr(config, "RETELL_AGENT_BUILDER_AGENT_ID", "agent_one")
    r = client.post(f"{BASE}/agents/{row['agent_id']}/test-call")
    assert r.status_code == 422
    assert r.json()["detail"]["leaked"]


def test_another_organizations_agent_is_not_there(client):
    aid = _draft(client)["done"]["agent_id"]
    row = store.load(aid)
    row["org_id"] = "some_other_org"
    store.save(row)
    assert client.get(f"{BASE}/agents/{aid}").status_code == 404
    assert aid not in [a["agent_id"] for a in client.get(f"{BASE}/agents").json()["agents"]]


def test_the_builder_is_not_on_the_public_surface():
    from services.security import matrix

    for route in matrix.routes(app):
        if route.path.startswith("/api/recruiter/agent-builder"):
            assert matrix.derive(route) == matrix.RECRUITER_AUTHENTICATED, route.path


def test_signed_out_is_refused(data_dir):
    assert TestClient(app).get(f"{BASE}/options").status_code == 401


def test_builder_open_demo_switch_opens_only_the_agent_builder(data_dir, monkeypatch):
    monkeypatch.setattr(config, "BUILDER_OPEN", True)
    monkeypatch.setattr(config, "ENVIRONMENT", "production")
    c = TestClient(app)
    assert c.get(f"{BASE}/options").status_code == 200
    # The rest of the recruiter console still needs a sign-in.
    assert c.get("/api/recruiter/interviews").status_code in (401, 503)


# --------------------------------------------------------------------------- #
#  The webhook
# --------------------------------------------------------------------------- #
def _sign(raw: bytes, key: str, ts: int) -> str:
    return f"v={ts},d=" + hmac.new(key.encode(), raw + str(ts).encode(), hashlib.sha256).hexdigest()


def test_signature_verification():
    raw, key, now = b'{"event":"call_ended"}', "k", int(time.time() * 1000)
    assert api.verify_signature(raw, _sign(raw, key, now), key, now)
    assert not api.verify_signature(raw + b" ", _sign(raw, key, now), key, now)
    assert not api.verify_signature(raw, _sign(raw, "other", now), key, now)
    assert not api.verify_signature(raw, _sign(raw, key, now - 6 * 60_000), key, now)
    assert not api.verify_signature(raw, "garbage", key, now)


def test_the_webhook_refuses_an_unsigned_body(data_dir, monkeypatch):
    monkeypatch.setattr(config, "RETELL_API_KEY", "k")
    r = TestClient(app).post("/api/agent-builder/retell-webhook", content=b"{}")
    assert r.status_code == 401


def test_the_webhook_ignores_calls_it_did_not_place(data_dir, monkeypatch):
    monkeypatch.setattr(config, "RETELL_API_KEY", "k")
    raw = json.dumps({"event": "call_ended", "call": {"call_id": "nobody", "metadata": {"source": "tara-agent-builder"}}}).encode()
    r = TestClient(app).post("/api/agent-builder/retell-webhook", content=raw,
                             headers={"x-retell-signature": _sign(raw, "k", int(time.time() * 1000))})
    assert r.status_code == 200 and r.json()["ignored"] is True


@pytest.mark.parametrize("word,want", [("Challenging", "Tough"), ("harder", "Realistic"), ("hard", "Tough"),
                                       ("easy", "Friendly"), ("Realistic", "Realistic")])
def test_difficulty_synonyms_land_on_a_real_value(word, want):
    # "harder" is not a known word, so the current value (Realistic) stays.
    assert ab._pick(word, rx.DIFFICULTIES, "Realistic") == want


@pytest.mark.parametrize("key", rx.VOICE_KEYS)
def test_the_persona_is_named_after_its_voice(key):
    p = ab.normalise_plan({"persona": {"name": "Dr. Emily Carter"}, "voice": key,
                           "opening_line": "Hi, I'm Emily Carter. Thanks for joining."})
    name = ab.voice_name(key)
    assert p["agent"]["persona"]["name"] == name
    assert "Emily" not in p["agent"]["opening_line"] and name in p["agent"]["opening_line"]


# --------------------------------------------------------------------------- #
#  The participant flow
# --------------------------------------------------------------------------- #
def _published_invite(client) -> tuple[str, str]:
    row = _draft(client)["done"]
    aid = row["agent_id"]
    client.put(f"{BASE}/agents/{aid}", json={**{k: row[k] for k in ("fields", "agent", "cfg")}, "reviewed": True})
    assert client.post(f"{BASE}/agents/{aid}/publish").status_code == 200
    inv = client.post(f"{BASE}/agents/{aid}/invites", json={"name": "Aarav", "email": "a@x.test"})
    assert inv.status_code == 200, inv.text
    assert inv.json()["link"].endswith("/participant?code=" + inv.json()["code"])
    return aid, inv.json()["code"]


def test_inviting_needs_a_published_agent(client):
    aid = _draft(client)["done"]["agent_id"]
    assert client.post(f"{BASE}/agents/{aid}/invites", json={}).status_code == 409


def test_a_participant_signs_in_and_sees_only_the_published_agent(client, data_dir):
    aid, code = _published_invite(client)
    p = TestClient(app)
    assert p.post("/api/participant/sign-in", json={"code": "RP-0000-ZZ", "name": "A", "email": "a@x.test", "consent": True}).json()["detail"]["error"] == "authFail"
    assert p.post("/api/participant/sign-in", json={"code": code, "name": "A", "email": "a@x.test", "consent": False}).status_code == 422
    r = p.post("/api/participant/sign-in", json={"code": code.lower(), "name": "Aarav Mehta", "email": "A@X.test", "consent": True})
    assert r.status_code == 200, r.text
    v = r.json()
    blob = json.dumps(v)
    agent = store.load(aid)["published"]["agent"]
    for r_ in agent["rubric"]:
        assert r_["name"] in v["agent"]["skills"]
        assert r_["anchor"] not in blob                  # descriptors stay hidden
    assert agent["instructions"] not in blob and agent["questions"][0]["text"] not in blob
    assert "weight" not in blob and "result" not in v and "evaluation" not in blob   # results are for admins only
    assert v["call"]["target_minutes"] in (10, 20, 30)
    sid = v["session_id"]
    assert p.get(f"/api/participant/session/{sid}").status_code == 200
    # another browser, no grant: the session does not exist
    assert TestClient(app).get(f"/api/participant/session/{sid}").status_code == 404


def test_the_call_carries_the_name_and_a_reconnect_resumes(client, monkeypatch):
    aid, code = _published_invite(client)
    p = TestClient(app)
    sid = p.post("/api/participant/sign-in", json={"code": code, "name": "Aarav Mehta", "email": "a@x.test", "consent": True}).json()["session_id"]
    monkeypatch.setattr(config, "RETELL_API_KEY", "key_test")
    monkeypatch.setattr(config, "RETELL_AGENT_BUILDER_AGENT_ID", "agent_one")
    _book_now(p, sid, monkeypatch)
    sent = []

    class R:
        def __init__(self, code, body):
            self.status_code, self._b = code, body

        def json(self):
            return self._b

    class FakeClient:
        def __init__(self, *a, **k):
            pass

        async def __aenter__(self):
            return self

        async def __aexit__(self, *a):
            return False

        async def post(self, url, json=None, headers=None):
            sent.append(json)
            return R(201, {"call_id": f"call_{len(sent)}", "access_token": "tok"})

        async def get(self, url, headers=None):
            return R(200, {"transcript_object": [{"role": "agent", "content": "Hi Aarav, tell me about RAG."},
                                                 {"role": "user", "content": "RAG retrieves documents at query time."}]})

    import httpx
    monkeypatch.setattr(httpx, "AsyncClient", FakeClient)
    first = p.post(f"/api/participant/session/{sid}/call", json={"speed": "slow"})
    assert first.status_code == 200 and first.json()["resumed"] is False
    v1 = sent[0]["retell_llm_dynamic_variables"]
    assert v1["candidate_name"] == "Aarav Mehta" and v1["resume_context"] == "none"
    assert sent[0]["agent_override"]["agent"]["voice_speed"] == 0.85
    assert sent[0]["metadata"]["participant_session"] == sid
    assert set(v1) == set(rx.VARIABLES)

    again = p.post(f"/api/participant/session/{sid}/call", json={})
    assert again.json()["resumed"] is True
    v2 = sent[1]["retell_llm_dynamic_variables"]
    assert "RAG retrieves documents" in v2["resume_context"]
    ov = sent[1]["agent_override"].get("conversation_flow") or sent[1]["agent_override"]["retell_llm"]
    assert ov["start_speaker"] == "agent"
    assert "cut off" in v2["opening_line"]


def _book_now(p, sid, monkeypatch):
    """Book the first offered slot and pretend it is that time."""
    from services.assessment import slots
    monkeypatch.setattr(slots, "retell_concurrency", lambda: (0, 20))
    start = p.get(f"/api/participant/session/{sid}/slots").json()["slots"][0]["start"]
    assert p.post(f"/api/participant/session/{sid}/booking", json={"start": start}).status_code == 200
    monkeypatch.setattr(slots, "can_join", lambda b, resuming=False, now=None: bool(b))


def test_complete_and_feedback_then_the_recruiter_sees_it(client):
    aid, code = _published_invite(client)
    p = TestClient(app)
    sid = p.post("/api/participant/sign-in", json={"code": code, "name": "Aarav", "email": "a@x.test", "consent": True}).json()["session_id"]
    assert p.post(f"/api/participant/session/{sid}/complete", json={"early": True, "elapsed_sec": 95}).status_code == 200
    assert p.post(f"/api/participant/session/{sid}/feedback", json={"rating": 4, "answers": {"clear": "Clear"}}).status_code == 200
    # submitted: the code cannot be used to sit it again
    again = p.post("/api/participant/sign-in", json={"code": code, "name": "Aarav", "email": "a@x.test", "consent": True})
    assert again.status_code == 409 and again.json()["detail"]["error"] == "completed"
    rows = client.get(f"{BASE}/agents/{aid}/sessions").json()["sessions"]
    assert rows[0]["status"] == "complete" and rows[0]["early"] is True and rows[0]["feedback"]["rating"] == 4


def test_every_agent_is_voice_only(client):
    row = _draft(client)["done"]
    aid = row["agent_id"]
    row["cfg"].update(format="Chat", image_proctoring=True)
    out = client.put(f"{BASE}/agents/{aid}", json={**{k: row[k] for k in ("fields", "agent", "cfg")}}).json()
    assert out["cfg"]["format"] == "Voice only"
    # Proctoring defaults are stored as chosen, for the proctoring suite.
    assert out["cfg"]["image_proctoring"] is True and out["cfg"]["safe_browser"] is False

def test_nothing_about_the_role_play_is_public_before_sign_in(client):
    _published_invite(client)
    assert TestClient(app).get("/api/participant/invite/RP-0000-ZZ").status_code in (404, 405)

def test_the_open_link_is_off_until_switched_on(client):
    row = _draft(client)["done"]
    aid = row["agent_id"]
    client.put(f"{BASE}/agents/{aid}", json={**{k: row[k] for k in ("fields", "agent", "cfg")}, "reviewed": True})
    pub = client.post(f"{BASE}/agents/{aid}/publish").json()
    assert pub["open_link"]["enabled"] is False and pub["candidate_link"] is None     # publishing doesn't open it
    on = client.post(f"{BASE}/agents/{aid}/open-link", json={"enabled": True}).json()
    assert on["enabled"] and on["link"].endswith(on["path"])
    code = on["code"]
    assert client.get(f"{BASE}/agents/{aid}").json()["open_link"]["code"] == code
    a, b = TestClient(app), TestClient(app)
    sa = a.post("/api/participant/sign-in", json={"code": code, "name": "A", "email": "a@x.test", "consent": True}).json()["session_id"]
    sb = b.post("/api/participant/sign-in", json={"code": code, "name": "B", "email": "b@x.test", "consent": True}).json()["session_id"]
    assert sa != sb                                                    # one session per person
    again = a.post("/api/participant/sign-in", json={"code": code, "name": "A", "email": "A@x.test", "consent": True}).json()["session_id"]
    assert again == sa                                                 # the same person rejoins
    # Switched off: the code stops working at once.
    assert client.post(f"{BASE}/agents/{aid}/open-link", json={"enabled": False}).json()["enabled"] is False
    c = TestClient(app)
    assert c.post("/api/participant/sign-in", json={"code": code, "name": "C", "email": "c@x.test", "consent": True}).status_code == 404
    # Back on: the same code works again.
    assert client.post(f"{BASE}/agents/{aid}/open-link", json={"enabled": True}).json()["code"] == code


def test_an_invitation_can_be_emailed_or_copied(client, monkeypatch):
    from services.notify import email

    aid, _ = _published_invite(client)
    r = client.post(f"{BASE}/agents/{aid}/invites", json={"name": "Maya Rao", "email": "maya@x.test", "send_email": True}).json()
    assert r["sent"] is False and "isn't set up" in r["email_error"]          # no SMTP here: nothing pretends to send
    assert r["code"] in r["text"] and r["link"] in r["text"] and r["text"].startswith("Hi Maya")
    sent = {}
    monkeypatch.setattr(email, "send", lambda to, subject, text, html=None: sent.update(to=to, subject=subject))
    r = client.post(f"{BASE}/agents/{aid}/invites", json={"name": "Lee", "email": "lee@x.test", "send_email": True, "message": "See you soon."}).json()
    assert r["sent"] is True and sent["to"] == "lee@x.test" and "See you soon." in r["text"]
    listed = client.get(f"{BASE}/agents/{aid}/invites").json()["invites"]
    assert {i["email"] for i in listed} >= {"maya@x.test", "lee@x.test"} and listed[0]["status"] == "Not started"
    bad = client.post(f"{BASE}/agents/{aid}/invites", json={"email": "not-an-email", "send_email": True})
    assert bad.status_code == 422


def test_every_call_carries_the_tuned_turn_taking():
    agent = rx.web_call_body(_row(), "agent_x")["agent_override"]["agent"]
    assert agent["responsiveness"] == 0.8 and agent["interruption_sensitivity"] == 0.7
    assert agent["backchannel_frequency"] == 0.2 and agent["enable_backchannel"] is True


def test_publishing_applies_the_current_rules_to_an_older_agent(client):
    row = _draft(client)["done"]
    aid = row["agent_id"]
    client.put(f"{BASE}/agents/{aid}", json={**{k: row[k] for k in ("fields", "agent", "cfg")}, "reviewed": True})
    stored = store.load(aid)
    stored["cfg"]["format"], stored["fields"]["type"] = "Video & voice", "Interview"   # saved under the old rules
    store.save(stored)
    pub = store.load(aid) if client.post(f"{BASE}/agents/{aid}/publish").status_code == 200 else None
    assert pub["published"]["cfg"]["format"] == "Voice only"
    assert pub["published"]["fields"]["type"] == "Role-play"



# --------------------------------------------------------------------------- #
#  Slots: never more live calls than Retell allows
# --------------------------------------------------------------------------- #
def _signed_in(client, n=1):
    aid, code = _published_invite(client)
    shared = client.post(f"{BASE}/agents/{aid}/open-link", json={"enabled": True}).json()["code"]
    out = []
    for i in range(n):
        c = TestClient(app)
        sid = c.post("/api/participant/sign-in", json={"code": shared, "name": f"P{i}", "email": f"p{i}@x.test", "consent": True}).json()["session_id"]
        out.append((c, sid))
    return out


def test_a_call_needs_a_booking_and_the_booked_time(client, monkeypatch):
    from services.assessment import slots
    monkeypatch.setattr(slots, "retell_concurrency", lambda: (0, 20))
    monkeypatch.setattr(config, "RETELL_API_KEY", "key_test")
    monkeypatch.setattr(config, "RETELL_AGENT_BUILDER_AGENT_ID", "agent_one")
    (p, sid), = _signed_in(client)
    r = p.post(f"/api/participant/session/{sid}/call", json={})
    assert r.status_code == 409 and r.json()["detail"]["error"] == "slot"          # no booking
    later = p.get(f"/api/participant/session/{sid}/slots").json()["slots"][-1]["start"]
    p.post(f"/api/participant/session/{sid}/booking", json={"start": later})
    r = p.post(f"/api/participant/session/{sid}/call", json={})
    assert r.status_code == 409 and r.json()["detail"]["error"] == "slot"          # booked, but not yet time


def test_a_slot_never_takes_more_people_than_the_limit(client, monkeypatch):
    monkeypatch.setenv("TARA_SLOT_CAPACITY", "2")
    people = _signed_in(client, 3)
    start = people[0][0].get(f"/api/participant/session/{people[0][1]}/slots").json()["slots"][0]["start"]
    codes = [c.post(f"/api/participant/session/{sid}/booking", json={"start": start}).status_code for c, sid in people]
    assert codes == [200, 200, 409]
    c, sid = people[2]
    left = {s["start"]: s["left"] for s in c.get(f"/api/participant/session/{sid}/slots").json()["slots"]}
    assert left[start] == 0
    # moving a booking frees its place
    c0, s0 = people[0]
    later = c0.get(f"/api/participant/session/{s0}/slots").json()["slots"][-1]["start"]
    assert c0.post(f"/api/participant/session/{s0}/booking", json={"start": later}).status_code == 200
    assert c.post(f"/api/participant/session/{sid}/booking", json={"start": start}).status_code == 200


def test_a_long_conversation_holds_every_block_it_runs_into():
    from datetime import datetime, timezone
    from services.assessment import slots
    t = datetime(2026, 10, 5, 10, 0, tzinfo=timezone.utc)
    assert len(slots.blocks_for(t, 25)) == 1 and len(slots.blocks_for(t, 35)) == 2 and len(slots.blocks_for(t, 60)) == 2


def test_the_joining_window():
    from datetime import datetime, timedelta, timezone
    from services.assessment import slots
    start = datetime(2026, 10, 5, 10, 0, tzinfo=timezone.utc)
    b = {"start": slots.iso(start), "cap_minutes": 25}
    assert not slots.can_join(b, now=start - timedelta(minutes=6))
    assert slots.can_join(b, now=start - timedelta(minutes=5))
    assert slots.can_join(b, now=start + timedelta(minutes=15))
    assert not slots.can_join(b, now=start + timedelta(minutes=16))
    assert slots.can_join(b, resuming=True, now=start + timedelta(minutes=30))     # a reconnect, still inside the booking


def test_no_call_when_retell_is_at_its_limit(client, monkeypatch):
    from services.assessment import slots
    monkeypatch.setattr(config, "RETELL_API_KEY", "key_test")
    monkeypatch.setattr(config, "RETELL_AGENT_BUILDER_AGENT_ID", "agent_one")
    (p, sid), = _signed_in(client)
    _book_now(p, sid, monkeypatch)
    monkeypatch.setattr(slots, "retell_concurrency", lambda: (20, 20))
    r = p.post(f"/api/participant/session/{sid}/call", json={})
    assert r.status_code == 503 and r.json()["detail"]["error"] == "busy"


def test_the_slot_under_way_can_still_be_booked_while_joinable():
    from datetime import datetime, timedelta, timezone
    from services.assessment import slots
    starts = slots.offered(now=datetime.now(timezone.utc))
    first = starts[0]
    assert first + timedelta(minutes=slots.JOIN_LATE_MIN) > datetime.now(timezone.utc)



def test_each_llm_use_has_its_own_model_slot():
    from services.ai.gateway import Workload, workload_config
    uses = {"agent_designer": "AGENT_DESIGNER_MODEL", "agent_reviser": "AGENT_REVISER_MODEL",
            "question_suggester": "QUESTION_SUGGESTER_MODEL", "rehearsal": "REHEARSAL_MODEL",
            "agent_scorer": "AGENT_SCORER_MODEL"}
    for w, var in uses.items():
        assert workload_config(Workload(w)).model == getattr(config, var)
    assert workload_config(Workload.AGENT_SCORER).temperature == 0.0


# --------------------------------------------------------------------------- #
#  Skill Master
# --------------------------------------------------------------------------- #
def test_the_skill_master_lists_readable_skills_for_every_domain(client):
    cat = client.get(f"{BASE}/skill-master").json()
    assert {d["category"] for d in cat["domains"]} == {"technical", "functional", "behavioural"}
    assert all(d["skills"] for d in cat["domains"])
    assert "Negotiation" in [s for d in cat["domains"] for s in d["skills"]]


def test_skills_added_from_the_master_get_anchors_and_questions(client, monkeypatch):
    row = _draft(client)["done"]
    aid = row["agent_id"]
    agent = copy.deepcopy(row["agent"])
    agent["rubric"].append({"name": "Negotiation", "anchor": "", "weight": 0})
    client.put(f"{BASE}/agents/{aid}", json={"fields": row["fields"], "agent": agent, "cfg": row["cfg"], "reviewed": True})
    monkeypatch.setattr(ab, "_complete", lambda *a, **k: {"skills": [
        {"name": "negotiation", "anchor": "Trades concessions for value.", "questions": ["What would you give up first?", "Why that price?"]},
        {"name": "Not in this agent", "anchor": "x", "questions": ["ignored"]}]})
    r = client.post(f"{BASE}/agents/{aid}/skills/draft", json={"names": ["Negotiation"]})
    assert r.status_code == 200, r.text
    out = r.json()
    neg = next(s for s in out["agent"]["rubric"] if s["name"] == "Negotiation")
    assert neg["anchor"] == "Trades concessions for value."
    assert [q["tag"] for q in out["agent"]["questions"]].count("Negotiation") == 2
    assert out["added"] == 2 and out["reviewed"] is False


def test_the_participant_never_hears_tara_only_the_persona():
    row = _row()
    row["agent"]["persona"]["name"] = "Dr. Maya Rao"
    row["agent"]["opening_line"] = "Hi, I'm Tara. Thanks for joining."
    row["agent"]["instructions"] = "Tara plays a skeptical buyer."
    v = rx.dynamic_variables(row, candidate_name="Tara Singh")
    assert "Tara" not in v["opening_line"] and "Maya" in v["opening_line"]
    assert v["conversation_instructions"].startswith("Maya plays")
    assert v["candidate_name"] == "Tara Singh"  # a participant may really be called Tara


def test_renaming_the_persona_reaches_every_line_and_question():
    agent = {"title": "Pricing call with Dr. Maya Rao", "opening_line": "Hi, I'm Maya Rao from Acme.",
             "closing_line": "Thanks, Maya signing off.", "instructions": "Maya is a skeptical buyer.",
             "description": "You'll speak with Dr. Maya Rao.", "questions": [{"text": "Maya asks why the price rose.", "tag": "x"}],
             "persona": {"name": "Dr. Maya Rao"}}
    ab.rename_persona(agent, "Dr. Maya Rao", "Adrian Cole")
    assert agent["title"] == "Pricing call with Adrian Cole"
    assert agent["opening_line"] == "Hi, I'm Adrian Cole from Acme."
    assert agent["closing_line"] == "Thanks, Adrian signing off."
    assert agent["instructions"].startswith("Adrian is")
    assert agent["questions"][0]["text"] == "Adrian asks why the price rose."
    ab.rename_persona(agent, "Adrian Cole", "")  # an empty name never wipes text
    assert "Adrian" in agent["opening_line"]


def test_the_config_never_keeps_tara_whoever_wrote_the_line(client):
    row = _draft(client)["done"]
    agent = copy.deepcopy(row["agent"])
    agent["persona"]["name"] = "Dr. Maya Rao"
    agent["instructions"] = "Tara conducts a friendly interview."
    agent["questions"][0]["text"] = "Tara asks about pricing."
    r = client.put(f"{BASE}/agents/{row['agent_id']}", json={"fields": row["fields"], "agent": agent, "cfg": row["cfg"], "reviewed": False})
    out = r.json()["agent"]
    assert out["instructions"] == "Maya conducts a friendly conversation."   # and never "interview"
    assert out["questions"][0]["text"] == "Maya asks about pricing."


def test_long_skill_names_are_not_a_leak_only_anchors_are():
    row = _row()
    row["agent"]["rubric"][0]["name"] = "Cross-Functional Collaboration"
    row["agent"]["questions"][0]["tag"] = "Cross-Functional Collaboration"
    row["fields"]["skills"] = "Cross-Functional Collaboration, design critique"
    assert rx.leaked_cues(rx.web_call_body(row, "agent_x"), row) == []


# --------------------------------------------------------------------------- #
#  Purpose, attempts, results
# --------------------------------------------------------------------------- #
def _publish_with(client, **cfg) -> tuple[str, str]:
    row = _draft(client)["done"]
    aid = row["agent_id"]
    client.put(f"{BASE}/agents/{aid}", json={"fields": row["fields"], "agent": row["agent"],
                                            "cfg": {**row["cfg"], **cfg}, "reviewed": True})
    assert client.post(f"{BASE}/agents/{aid}/publish").status_code == 200
    inv = client.post(f"{BASE}/agents/{aid}/invites", json={"name": "Lee", "email": "l@x.test"})
    return aid, inv.json()["code"]


def test_there_is_no_purpose_field_and_attempts_default_to_one(client):
    row = _draft(client)["done"]   # even "a technical interview" brief: no purpose is guessed or stored
    assert "purpose" not in row["cfg"] and row["cfg"]["attempts"] == "1"
    out = client.put(f"{BASE}/agents/{row['agent_id']}", json={"fields": row["fields"], "agent": row["agent"],
                     "cfg": {**row["cfg"], "purpose": "L&D", "attempts": "Unlimited"}}).json()
    assert "purpose" not in out["cfg"] and out["cfg"]["attempts"] == "Unlimited"

def test_a_learner_can_practise_again_but_never_sees_a_result(client, monkeypatch):
    from services.api import participant as part

    aid, code = _publish_with(client, purpose="L&D", attempts="Unlimited")
    p = TestClient(app)
    s1 = p.post("/api/participant/sign-in", json={"code": code, "name": "Lee", "email": "l@x.test", "consent": True}).json()
    assert s1["attempt"] == 1 and s1["purpose"] == "General"
    done = p.post(f"/api/participant/session/{s1['session_id']}/complete", json={"early": False, "elapsed_sec": 300}).json()
    assert "result" not in done and "result_status" not in done      # admins only, even for learners
    s2 = p.post("/api/participant/sign-in", json={"code": code, "name": "Lee", "email": "l@x.test", "consent": True}).json()
    assert s2["attempt"] == 2 and s2["session_id"] != s1["session_id"]


def test_a_limited_attempt_policy_is_enforced(client):
    aid, code = _publish_with(client, purpose="L&D", attempts="3")
    p = TestClient(app)
    for n in range(3):
        sid = p.post("/api/participant/sign-in", json={"code": code, "name": "Lee", "email": "l@x.test", "consent": True}).json()["session_id"]
        p.post(f"/api/participant/session/{sid}/complete", json={"early": False, "elapsed_sec": 60})
    r = p.post("/api/participant/sign-in", json={"code": code, "name": "Lee", "email": "l@x.test", "consent": True})
    assert r.status_code == 409 and "all 3 attempts" in r.json()["detail"]["message"]


def test_a_session_is_evaluated_once_and_only_admins_see_it(client, monkeypatch):
    from services.api import participant as part
    from services.data import agent_sessions
    from tests.test_simulation_evaluation import EVIDENCE, JUDGED, TRANSCRIPT, FakeLLM

    aid, code = _published_invite(client)                     # a Hiring agent
    p = TestClient(app)
    sid = p.post("/api/participant/sign-in", json={"code": code, "name": "A", "email": "a@x.test", "consent": True}).json()["session_id"]
    row = agent_sessions.load(sid)
    row["snapshot"]["agent"]["rubric"] = [{"name": "Negotiation", "anchor": "x", "weight": 60},
                                          {"name": "Discovery", "anchor": "y", "weight": 40}]
    row.update(status="complete", transcript=TRANSCRIPT, calls=["call_1"])
    llm = FakeLLM(EVIDENCE, JUDGED)
    monkeypatch.setattr(ab, "_complete", llm)
    part.evaluate_session(row)
    part.evaluate_session(row)                                # a retried webhook / second completion
    assert row["evaluation"]["purpose"] == "General" and len(llm.calls) == 3
    agent_sessions.save(row)
    v = p.get(f"/api/participant/session/{sid}").json()
    assert "result" not in v and "evaluation" not in json.dumps(v)
    recruiter = client.get(f"{BASE}/agents/{aid}/sessions").json()["sessions"][0]
    assert recruiter["evaluation"]["overall"] is not None


def test_proctoring_is_fields_for_the_suite_and_nothing_else(client):
    aid, _ = _publish_with(client, purpose="Hiring", image_proctoring=True)
    inv = client.get(f"{BASE}/agents/{aid}/invitation").json()
    assert inv["defaults"] == {"image_proctoring": True, "safe_browser": False} and inv["label"] == "AI Conversation"
    r = client.post(f"{BASE}/agents/{aid}/invitations", json={"emails": ["a@x.test"], "image_proctoring": False, "safe_browser": True}).json()
    code = r["results"][0]["code"]
    v = TestClient(app).post("/api/participant/sign-in", json={"code": code, "name": "Lee", "email": "a@x.test", "consent": True}).json()
    # The invitation's settings travel with the session for the suite; the page and the call don't act on them.
    assert v["proctoring"] == {"image_proctoring": False, "safe_browser": True} and v["call"]["camera_required"] is False
    row = store.load(aid)
    body = rx.web_call_body(row["published"] | {"agent_id": aid}, "agent_x")
    assert "proctor" not in json.dumps(body).lower() and "safe_browser" not in json.dumps(body)
    rep = client.get(f"{BASE}/agents/{aid}/report").json()
    assert any(x["proctoring"] == "Safe browser" for x in rep["rows"])

def test_a_published_role_play_is_locked_and_can_be_duplicated(client, monkeypatch):
    aid, code = _published_invite(client)
    row = client.get(f"{BASE}/agents/{aid}").json()
    assert row["locked"] is True
    body = {k: row[k] for k in ("fields", "agent", "cfg")}
    for method, path, payload in [("put", f"/agents/{aid}", body), ("post", f"/agents/{aid}/publish", None),
                                  ("post", f"/agents/{aid}/revise", {"instruction": "tougher"}),
                                  ("post", f"/agents/{aid}/questions/generate", None),
                                  ("post", f"/agents/{aid}/skills/draft", {"names": ["x"]})]:
        r = getattr(client, method)(BASE + path, json=payload) if payload is not None else getattr(client, method)(BASE + path)
        assert r.status_code == 409 and r.json()["detail"]["error"] == "locked", path
    # Still usable: invite, open link, results.
    assert client.post(f"{BASE}/agents/{aid}/invites", json={"name": "B"}).status_code == 200
    assert client.post(f"{BASE}/agents/{aid}/open-link", json={"enabled": True}).status_code == 200
    assert client.get(f"{BASE}/agents/{aid}/sessions").status_code == 200
    # Duplicate: an editable draft with none of the original's people or results.
    dup = client.post(f"{BASE}/agents/{aid}/duplicate").json()
    assert dup["agent_id"] != aid and dup["locked"] is False and dup["published_version"] == 0
    assert dup["agent"]["title"].endswith("(copy)") and dup["open_link"]["enabled"] is False
    assert client.put(f"{BASE}/agents/{dup['agent_id']}", json={k: dup[k] for k in ("fields", "agent", "cfg")}).status_code == 200


def test_the_report_grid_and_an_admin_review(client, monkeypatch):
    from services.api import participant as part
    from services.data import agent_sessions
    from tests.test_simulation_evaluation import EVIDENCE, JUDGED, TRANSCRIPT, FakeLLM

    aid, code = _published_invite(client)
    p = TestClient(app)
    sid = p.post("/api/participant/sign-in", json={"code": code, "name": "Asha", "email": "asha@x.test", "consent": True}).json()["session_id"]
    s = agent_sessions.load(sid)
    s["snapshot"]["agent"]["rubric"] = [{"name": r["name"], "anchor": "x", "weight": r["weight"]} for r in s["snapshot"]["agent"]["rubric"]]
    first = s["snapshot"]["agent"]["rubric"][0]["name"]
    ev = [dict(e, skill=first) for e in EVIDENCE]
    s.update(status="complete", transcript=TRANSCRIPT, calls=["c1"], elapsed_sec=310)
    monkeypatch.setattr(ab, "_complete", FakeLLM(ev, [dict(JUDGED[0], name=first)]))
    part.evaluate_session(s)
    agent_sessions.save(s)
    client.post(f"{BASE}/agents/{aid}/invites", json={"name": "Not Yet", "email": "later@x.test"})
    rep = client.get(f"{BASE}/agents/{aid}/report").json()
    assert rep["agent"]["purpose"] == "General" and rep["decisions"] == ["Recommended", "Needs review", "Not recommended"]
    row = next(x for x in rep["rows"] if x["session_id"] == sid)
    assert row["name"] == "Asha" and row["status"] == "Completed" and row["evaluation"] == "Evaluated" and row["duration_sec"] == 310
    # The invitation that was never used is in the grid as Pending.
    assert any(x["status"] == "Pending" for x in rep["rows"])
    assert row["skills"][first] is not None and rep["stats"]["evaluated"] == 1 and rep["stats"]["average"] == row["overall"]
    assert client.put(f"{BASE}/agents/{aid}/attempts/{sid}/review", json={"decision": "Maybe"}).status_code == 422
    r = client.put(f"{BASE}/agents/{aid}/attempts/{sid}/review", json={"decision": "Recommended", "notes": "Strong discovery."}).json()
    assert r["decision"] == "Recommended" and r["by"]
    again = client.get(f"{BASE}/agents/{aid}/report").json()
    mine = next(x for x in again["rows"] if x["session_id"] == sid)
    assert mine["review"]["notes"] == "Strong discovery." and again["stats"]["decisions"]["Recommended"] == 1
    assert mine["recommendation"] == "Recommended" and mine["recommendation_source"] == "admin" and mine["recommendation_level"] == "positive"
    # Org-wide reports, and Excel downloads that respect the filters.
    allr = client.get(f"{BASE}/reports").json()
    assert any(x["session_id"] == sid and x["agent_title"] for x in allr["rows"]) and allr["agents"]
    x = client.get(f"{BASE}/agents/{aid}/report.xlsx", params={"status": "Completed"})
    assert x.status_code == 200 and x.content[:2] == b"PK" and "spreadsheetml" in x.headers["content-type"]
    import io
    from openpyxl import load_workbook
    ws = load_workbook(io.BytesIO(x.content)).active
    assert ws.max_row == 2 and ws["A2"].value == "Asha" and ws["K2"].value == "Recommended"
    assert client.get(f"{BASE}/reports.xlsx").status_code == 200
    detail = client.get(f"{BASE}/agents/{aid}/attempts/{sid}/report").json()
    assert detail["evaluation"]["overall"] == row["overall"] and detail["transcript"]
    # The AI's recommendation is kept beside the admin's decision, never replaced.
    assert detail["evaluation"]["recommendation"] and detail["review"]["decision"] == "Recommended"
    # Another agent's attempt is not reachable through this agent.
    other = client.post(f"{BASE}/agents/{aid}/duplicate").json()["agent_id"]
    assert client.get(f"{BASE}/agents/{other}/attempts/{sid}/report").status_code == 404


def test_bulk_invitations_use_a_safe_template(client, monkeypatch):
    from services.notify import email, template as tmpl

    aid, _ = _published_invite(client)
    bad = client.post(f"{BASE}/agents/{aid}/invitations", json={"emails": ["ok@x.test", "nope"]})
    assert bad.status_code == 422 and "nope" in bad.json()["detail"]
    assert client.post(f"{BASE}/agents/{aid}/invitations", json={"emails": [f"p{i}@x.test" for i in range(11)]}).status_code == 422
    sent = []
    monkeypatch.setattr(email, "configured", lambda: True)
    monkeypatch.setattr(email, "send", lambda to, subject, text, html=None: sent.append((to, html)))
    tpl = '<p>Hi {PARTICIPANT_NAME}, welcome to <b>{ROLE_PLAY}</b>.</p><script>alert(1)</script><img src=x onerror=alert(1)><a href="javascript:x">x</a>'
    r = client.post(f"{BASE}/agents/{aid}/invitations", json={"emails": ["maya.rao@x.test", "MAYA.RAO@x.test", "lee@x.test"], "template_html": tpl}).json()
    assert [x["email"] for x in r["results"]] == ["maya.rao@x.test", "lee@x.test"]            # de-duplicated
    assert all(x["sent"] for x in r["results"]) and len(sent) == 2
    html_ = sent[0][1]
    assert "Hi Maya," in html_ and "<b>" in html_ and r["results"][0]["code"] in html_ and r["results"][0]["link"] in html_
    assert "<script" not in html_ and "onerror" not in html_ and "javascript:" not in html_ and "<img" not in html_
    assert tmpl.sanitize('<div onclick="x">a<style>b</style></div>') == "<p>a</p>"


def test_the_word_interview_never_reaches_anyone():
    assert api.no_interview("Senior AI Engineer Interview: two interviews, an interview. An interview.") == \
        "Senior AI Engineer Conversation: two conversations, a conversation. A conversation."
    out = api.no_interview("Interview prep: this interview has interviews.")
    assert "nterview" not in out and out.startswith("Conversation prep")
    agent = api.clean_agent({"title": "Data Scientist Interview", "type_label": "Technical interview",
                             "persona": {"name": "Maya"}, "description": "A short interview.",
                             "instructions": "Run the interview.", "opening_line": "Welcome to the interview.",
                             "closing_line": "Thanks.", "questions": [{"text": "Why this interview?", "tag": "x"}]})
    assert "nterview" not in json.dumps(agent)


# --------------------------------------------------------------------------- #
#  Exhibits: charts and images the participant sees during the conversation
# --------------------------------------------------------------------------- #
PNG = (b"\x89PNG\r\n\x1a\n\x00\x00\x00\rIHDR\x00\x00\x00\x01\x00\x00\x00\x01\x08\x06\x00\x00\x00\x1f\x15\xc4\x89"
       b"\x00\x00\x00\rIDATx\x9cc\xf8\x0f\x00\x00\x01\x01\x00\x05\x18\xd8N\x00\x00\x00\x00IEND\xaeB`\x82")


def test_exhibits_reach_the_persona_and_the_participant_only(client):
    import base64

    row = _draft(client)["done"]
    aid = row["agent_id"]
    bad = client.post(f"{BASE}/agents/{aid}/exhibit-images", json={"data_base64": base64.b64encode(b"<svg onload=x>hi</svg>").decode()})
    assert bad.status_code == 422                                              # checked by its bytes
    up = client.post(f"{BASE}/agents/{aid}/exhibit-images", json={"data_base64": "data:image/png;base64," + base64.b64encode(PNG).decode()})
    assert up.status_code == 200, up.text
    f = up.json()["file"]
    agent = copy.deepcopy(row["agent"])
    agent["exhibits"] = [
        {"kind": "chart", "title": "Quarterly revenue", "description": "Q3 dipped after a price change.",
         "chart": {"type": "bar", "labels": ["Q1", "Q2", "Q3"], "values": [120, 140, "95"], "unit": "k"}},
        {"kind": "image", "title": "Pricing sheet", "description": "List prices for the three plans.", "file": f},
        {"kind": "image", "title": "Bad", "file": "../../etc/passwd"},             # dropped: not an uploaded file
    ]
    out = client.put(f"{BASE}/agents/{aid}", json={"fields": row["fields"], "agent": agent, "cfg": row["cfg"], "reviewed": True}).json()
    assert [e["title"] for e in out["agent"]["exhibits"]] == ["Quarterly revenue", "Pricing sheet"]
    assert out["agent"]["exhibits"][0]["chart"]["values"] == [120.0, 140.0, 95.0]
    v = rx.dynamic_variables(store.load(aid))["conversation_instructions"]
    assert "Exhibit 1, Quarterly revenue (bar chart: Q1: 120 k, Q2: 140 k, Q3: 95 k)" in v and "Exhibit 2, Pricing sheet" in v
    assert client.get(f"{BASE}/agents/{aid}/exhibit-images/{f}").status_code == 200
    assert client.get(f"{BASE}/agents/{aid}/exhibit-images/{'0' * 24}.png").status_code == 404
    assert client.post(f"{BASE}/agents/{aid}/publish").status_code == 200
    inv = client.post(f"{BASE}/agents/{aid}/invites", json={"name": "P"}).json()["code"]
    p = TestClient(app)
    s = p.post("/api/participant/sign-in", json={"code": inv, "name": "P", "email": "p@x.test", "consent": True}).json()
    ex = s["agent"]["exhibits"]
    assert ex[0]["n"] == 1 and ex[0]["chart"]["type"] == "bar" and ex[1]["image"].endswith(f)
    img = p.get(ex[1]["image"])
    assert img.status_code == 200 and img.headers["content-type"] == "image/png" and img.content == PNG
    assert TestClient(app).get(ex[1]["image"]).status_code == 404                # another browser: no session, no image
    assert "description" not in json.dumps(ex)                                  # what the persona knows stays with the persona


def test_capacity_is_18_by_default_and_never_above_retell(monkeypatch):
    from services.assessment import slots

    monkeypatch.delenv("TARA_SLOT_CAPACITY", raising=False)
    monkeypatch.setattr(slots, "_cap_cache", (0.0, 0))
    monkeypatch.setattr(slots, "retell_concurrency", lambda: (0, 20))
    assert slots.capacity() == 18
    monkeypatch.setenv("TARA_SLOT_CAPACITY", "25")
    assert slots.capacity() == 20                                   # Retell's limit wins when lower
    monkeypatch.setattr(slots, "_cap_cache", (0.0, 0))
    monkeypatch.setattr(slots, "retell_concurrency", lambda: (0, 0))  # Retell can't be asked
    assert slots.capacity() == 25


def test_start_now_is_offered_only_when_there_is_room(client, monkeypatch):
    from services.assessment import slots

    monkeypatch.setenv("TARA_SLOT_CAPACITY", "1")
    monkeypatch.setattr(slots, "_cap_cache", (0.0, 0))
    monkeypatch.setattr(slots, "retell_concurrency", lambda: (0, 20))
    (p1, s1), (p2, s2) = _signed_in(client, 2)
    now = p1.get(f"/api/participant/session/{s1}/slots").json()["now"]
    assert now and now["now"] is True and now["left"] == 1
    b = p1.post(f"/api/participant/session/{s1}/booking", json={"start": "now"})
    assert b.status_code == 200 and b.json()["booking"]["can_join_now"] is True
    # The only place right now is taken: the second person isn't offered "now" and can't book it.
    assert p2.get(f"/api/participant/session/{s2}/slots").json()["now"] is None
    full = p2.post(f"/api/participant/session/{s2}/booking", json={"start": "now"})
    assert full.status_code == 409 and full.json()["detail"]["error"] == "full"
    # Live calls at the cap also close "now", whatever the bookings say.
    monkeypatch.setattr(slots, "retell_concurrency", lambda: (1, 20))
    assert slots.now_option([], 25) is None


def test_start_now_respects_the_daily_hours(monkeypatch):
    from datetime import datetime, timezone

    from services.assessment import slots

    monkeypatch.setattr(slots, "HOURS", "09:00-21:00")
    monkeypatch.setattr(slots, "TZ", "Asia/Kolkata")
    monkeypatch.setattr(slots, "retell_concurrency", lambda: (0, 20))
    monkeypatch.setattr(slots, "_cap_cache", (0.0, 0))
    assert slots.now_option([], 25, now=datetime(2026, 10, 3, 5, 0, tzinfo=timezone.utc)) is not None    # 10:30 IST
    assert slots.now_option([], 25, now=datetime(2026, 10, 3, 18, 0, tzinfo=timezone.utc)) is None       # 23:30 IST


def test_slots_are_open_around_the_clock_by_default(monkeypatch):
    from datetime import datetime, timezone

    from services.assessment import slots

    monkeypatch.setattr(slots, "HOURS", "00:00-24:00")
    monkeypatch.setattr(slots, "retell_concurrency", lambda: (0, 20))
    monkeypatch.setattr(slots, "_cap_cache", (0.0, 0))
    at = datetime(2026, 10, 3, 20, 45, tzinfo=timezone.utc)                       # 02:15 IST
    assert slots.now_option([], 25, now=at) is not None
    starts = slots.offered(at)
    assert len({t.astimezone(slots.ZoneInfo(slots.TZ)).strftime("%H:%M") for t in starts}) == 48   # every half hour


# --------------------------------------------------------------------------- #
#  Retell v3 browser client, pointed at our server
# --------------------------------------------------------------------------- #
def test_the_v3_relay_serves_only_your_own_call(client, monkeypatch):
    import contextlib

    from services.api import retell_relay
    from services.data import agent_sessions
    from starlette.websockets import WebSocketDisconnect

    (p, sid), = _signed_in(client)
    row = agent_sessions.load(sid)
    row["calls"] = ["call_mine"]
    agent_sessions.save(row)
    stopped = []

    async def fake_stop(cid):
        stopped.append(cid)
    monkeypatch.setattr(retell_relay, "_stop", fake_stop)
    base = f"/api/participant/session/{sid}/retell"
    assert p.post(f"{base}/v2/stop-call/call_mine").status_code == 200 and stopped == ["call_mine"]
    assert p.post(f"{base}/v2/stop-call/call_someone_else").status_code == 404     # not this participant's call
    assert TestClient(app).post(f"{base}/v2/stop-call/call_mine").status_code == 404  # no session cookie

    class FakeUpstream:
        close_code = 1000
        def __init__(self):
            self.msgs = ['{"type":"transcript_updated","transcript":[{"id":"1","role":"agent","content":"Hello"}]}']
        def __aiter__(self):
            return self
        async def __anext__(self):
            if not self.msgs:
                raise StopAsyncIteration
            return self.msgs.pop(0)
        async def send(self, m):
            pass

    @contextlib.asynccontextmanager
    async def fake_connect(url, subprotocols=None, open_timeout=None):
        assert url.endswith("/v2/monitor-call/call_mine") and subprotocols[0] == "bearer"
        yield FakeUpstream()

    import websockets
    monkeypatch.setattr(websockets, "connect", fake_connect)
    with p.websocket_connect(f"{base}/v2/monitor-call/call_mine", subprotocols=["bearer", "session"]) as ws:
        assert "Hello" in ws.receive_text()
    with pytest.raises(WebSocketDisconnect):
        with p.websocket_connect(f"{base}/v2/monitor-call/call_someone_else", subprotocols=["bearer", "session"]) as ws:
            ws.receive_text()


def test_assessment_is_on_hold_everywhere(client):
    d = _draft(client, brief="An assessment of SQL skills for analysts.", mode="assessment")["done"]
    assert d["fields"]["type"] == "Role-play" and d["mode"] == "roleplay"
    out = client.put(f"{BASE}/agents/{d['agent_id']}", json={"fields": {**d["fields"], "type": "Assessment"},
                     "agent": d["agent"], "cfg": d["cfg"]}).json()
    assert out["fields"]["type"] == "Role-play"


def test_a_fast_rejoin_uses_the_pages_transcript_and_the_real_time_used(client, monkeypatch):
    aid, code = _published_invite(client)
    p = TestClient(app)
    sid = p.post("/api/participant/sign-in", json={"code": code, "name": "Aarav", "email": "a@x.test", "consent": True}).json()["session_id"]
    monkeypatch.setattr(config, "RETELL_API_KEY", "key_test")
    monkeypatch.setattr(config, "RETELL_AGENT_BUILDER_AGENT_ID", "agent_one")
    monkeypatch.setattr(ab, "_complete", lambda *a, **k: (_ for _ in ()).throw(RuntimeError("no model in tests")))
    _book_now(p, sid, monkeypatch)
    sent = []

    class R:
        def __init__(self, code, body): self.status_code, self._b = code, body
        def json(self): return self._b

    class FakeClient:
        def __init__(self, *a, **k): pass
        async def __aenter__(self): return self
        async def __aexit__(self, *a): return False
        async def post(self, url, json=None, headers=None):
            sent.append(json)
            return R(201, {"call_id": f"call_{len(sent)}", "access_token": "tok"})
        async def get(self, url, headers=None):
            # Retell hasn't finished the dropped call's transcript, but knows it lasted 9 minutes.
            return R(200, {"transcript_object": [], "start_timestamp": 0, "end_timestamp": 9 * 60_000})

    import httpx
    monkeypatch.setattr(httpx, "AsyncClient", FakeClient)
    assert p.post(f"/api/participant/session/{sid}/call", json={}).status_code == 200
    page = [{"role": "agent", "text": "Walk me through how you'd evaluate a RAG system."},
            {"role": "user", "text": "First I'd build a labelled set of questions and"}]
    again = p.post(f"/api/participant/session/{sid}/call", json={"transcript": page, "elapsed_sec": 300}).json()
    assert again["resumed"] is True
    v = sent[1]["retell_llm_dynamic_variables"]
    assert "labelled set of questions" in v["resume_context"] and "may not have finished" in v["resume_context"]
    assert "About 9 minutes are already used" in v["resume_context"]           # Retell's time, not the page's
    assert "middle of your answer" in v["opening_line"]                        # no model: the code's rejoin line


def test_audio_setup_is_passed_through_to_retell_only_for_your_own_call(client, monkeypatch):
    import httpx

    from services.data import agent_sessions

    (p, sid), = _signed_in(client)
    row = agent_sessions.load(sid)
    row["calls"] = ["call_mine"]
    agent_sessions.save(row)
    seen = {}

    class R:
        status_code = 201
        content = b'{"session_id":"s1","sdp":"answer"}'
        headers = {"content-type": "application/json"}

    class FakeClient:
        def __init__(self, *a, **k): pass
        async def __aenter__(self): return self
        async def __aexit__(self, *a): return False
        async def request(self, method, url, headers=None, content=None):
            seen.update(method=method, url=url, auth=headers.get("authorization"), body=content)
            return R()

    monkeypatch.setattr(httpx, "AsyncClient", FakeClient)
    base = f"/api/participant/session/{sid}/retell/webrtc-proxy"
    r = p.post(f"{base}/call_mine/v1/webrtc/sessions", headers={"Authorization": "Bearer call_token"}, json={"sdp": "offer"})
    assert r.status_code == 201 and r.json()["session_id"] == "s1"
    assert seen["url"] == "https://api.retellai.com/webrtc-proxy/call_mine/v1/webrtc/sessions"
    assert seen["auth"] == "Bearer call_token" and b"offer" in seen["body"]      # the call's own token, as sent
    assert p.post(f"{base}/call_other/v1/webrtc/sessions", json={}).status_code == 404   # not your call
    assert p.post(f"{base}/call_mine/v2/anything-else", json={}).status_code == 404      # only WebRTC signalling


# --------------------------------------------------------------------------- #
#  The page itself sits behind the sign-in
# --------------------------------------------------------------------------- #
def test_the_builder_page_shows_the_sign_in_until_there_is_a_session(data_dir, tenant):
    anonymous = TestClient(app)
    first = anonymous.get("/agent-builder")
    assert first.status_code == 200
    assert 'id="f"' in first.text and "/api/auth/login" in first.text
    assert "no-store" in first.headers["cache-control"]

    sign_in(anonymous, tenant)
    after = anonymous.get("/agent-builder")
    assert 'id="f"' not in after.text and "/api/auth/login" not in after.text


def test_builder_open_does_not_skip_the_page_sign_in(data_dir, tenant, monkeypatch):
    monkeypatch.setattr(config, "BUILDER_OPEN", True)
    page = TestClient(app).get("/agent-builder")
    assert 'id="f"' in page.text
